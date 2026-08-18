"""
Build v2.5 from v2.4: preserves all 228 existing sheets, appends the
UOC-COMPILER-FULL-RECONFIGURATION-001 sheets.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import json

SRC = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v2.4.xlsx"
OUT = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v2.5.xlsx"
RECON = "/home/user/URSP/reconstruction"

def load(name):
    with open(f"{RECON}/{name}") as f:
        return json.load(f)

KERNEL = load("uoc_compiler_kernel.json")
IR = load("uoc_compiler_ir.json")
PASSES = load("uoc_compiler_passes.json")
INV = load("uoc_compiler_invariants.json")
EQUIV = load("uoc_compiler_equivalences.json")
FIXED = load("uoc_compiler_fixed_points.json")
MDCL = load("uoc_master_mdcl.json")
CLOSURE = load("uoc_master_closure_matrix.json")
PROOF = load("uoc_master_proof_registry.json")
FALS = load("uoc_master_falsification_ledger.json")
HIST = load("uoc_historical_reclassification.json")

wb = openpyxl.load_workbook(SRC, data_only=False)

TITLE_FONT = Font(name="Arial", size=12, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")

def new_sheet(name):
    return wb.create_sheet(name)

def write_title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = WRAP

def write_table(ws, start_row, headers, rows, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + 1 + len(rows)

def block(ws, r, title, text, height=60):
    ws.cell(row=r, column=1, value=title).font = Font(bold=True)
    r += 1
    ws.cell(row=r, column=1, value=text)
    ws.cell(row=r, column=1).font = BODY_FONT
    ws.cell(row=r, column=1).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = height
    return r + 2


# 229 Compiler Kernel
ws = new_sheet("229 RECFG-001 Kernel")
write_title(ws, "COMPILER KERNEL RECONFIGURATION", "")
r = block(ws, 5, "Diagnosis", KERNEL["the_reconfiguration"]["diagnosis"], height=110)
r = block(ws, r, "Resolution", KERNEL["the_reconfiguration"]["resolution"], height=90)
r = block(ws, r, "Theorem proven this run", KERNEL["the_reconfiguration"]["theorem_proven_this_run"], height=110)
r = block(ws, r, "Honest residual finding", KERNEL["the_reconfiguration"]["honest_residual_finding"], height=90)

# 230 Kernel Candidates
ws = new_sheet("230 RECFG-002 Candidates")
write_title(ws, "KERNEL CANDIDATES", "")
write_table(ws, 5, ["Candidate", "K bits (N=4)"], [
    ("R", KERNEL["complexity_comparison_at_N_4"]["K_R_bits"]),
    ("R + derived N_orient", KERNEL["complexity_comparison_at_N_4"]["K_R_plus_derived_N_bits"]),
    ("(D,T,C) independent, worst case", KERNEL["complexity_comparison_at_N_4"]["K_DTC_independent_worst_case_bits"]),
], widths=[35, 60])
r = 10
r = block(ws, r, "Kernel definition K*", KERNEL["kernel_definition"]["K_star"], height=70)
r = block(ws, r, "Is the kernel single-valued?", KERNEL["kernel_definition"]["is_the_kernel_single_valued"], height=90)

# 231 Seed Reconfiguration
ws = new_sheet("231 RECFG-003 Seed Reconfig")
write_title(ws, "SEED RECONFIGURATION -- CORRECTED F_N^derived", "")
counts = FIXED["F_N_derived_v2"]["counts"]
write_table(ws, 5, ["N", "|F_N^derived_v2|"], [(k, v) for k, v in counts.items()], widths=[10, 20])
r = 10
r = block(ws, r, "All rigid?", FIXED["F_N_derived_v2"]["all_rigid"], height=30)
r = block(ws, r, "Symmetric overlap", FIXED["F_N_derived_v2"]["symmetric_overlap"], height=70)
r = block(ws, r, "Closure object: K* or family?", FIXED["is_the_closure_object_K_star_or_a_family"]["conclusion"], height=90)

# 232 Relation Audit
ws = new_sheet("232 RECFG-004 Relation Audit")
write_title(ws, "RELATION / EQUIVALENCE AUDIT", "")
r = block(ws, 5, "Correction to prior notation", EQUIV["correction"], height=90)
rows = [(k, v) for k, v in EQUIV["fixed_point_taxonomy"].items()]
write_table(ws, r, ["Notion", "Definition/status"], rows, widths=[30, 110])

# 233 Spectral Reconfiguration
ws = new_sheet("233 RECFG-005 Spectral")
write_title(ws, "SPECTRAL RECONFIGURATION", "")
r = block(ws, 5, "Resolution search outcome", "The candidate list (R+R^T, |A_R|, R^TR, Hermitian/magnetic adjacency, directed Laplacian, etc.) reduces, once the kernel is correctly reconfigured, to the simplest option already in use: A_sym=(A+A^T)/2 -- the seed is now free to be symmetric in principle, so no exotic directed-spectral extension was shown necessary. Not separately implemented/tested this run since the simpler fix already resolves the specific incompatibility.", height=120)
r = block(ws, r, "Operators kept distinct", "L=D-A_sym; Q=-D^-1 L; Q*=-L D^-1 -- THM-C-004D-1 (Q != -L for irregular graphs) reaffirmed, not re-litigated.", height=50)

# 234 ARBS Placement
ws = new_sheet("234 RECFG-006 ARBS Placement")
write_title(ws, "ARBS PLACEMENT IN THE COMPILER", "")
r = block(ws, 5, "Determination", "ARBS supplies exactly two concrete, checkable admissibility axioms (TH-ARBS-001A/B) used as FILTERS on the C0 kernel candidates -- intermediate compiler IR / admissibility metadata, NOT the primitive layer. No concrete G_ARBS=(V,E,W) instantiation exists anywhere in the corpus (unchanged from ARBS-GRAPH-REALIZATION-001) to BE the primitive layer.", height=110)
r = block(ws, r, "tilde_G_k placement", "Independent validation substrate, admissible member of Graph_ARBS, downstream of (not identical to) the C0 kernel layer.", height=60)

# 235 UOC-IR
ws = new_sheet("235 RECFG-007 UOC-IR")
write_title(ws, "UOC-IR SCHEMA", "")
rows = [(k, v) for k, v in IR["UOC_IR_schema"].items()]
write_table(ws, 5, ["Field", "Definition"], rows, widths=[22, 120])
r = 18
r = block(ws, r, "Self-representability", IR["self_representability_test"]["result"], height=100)

# 236 Compiler Passes
ws = new_sheet("236 RECFG-008 Passes")
write_title(ws, "COMPILER PASSES", "")
rows = [(p["id"], p["name"], p["status"], p["detail"]) for p in PASSES["passes"]]
write_table(ws, 5, ["ID", "Name", "Status", "Detail"], rows, widths=[14, 22, 26, 90])

# 237 Compiler Invariants
ws = new_sheet("237 RECFG-009 Invariants")
write_title(ws, "COMPILER INVARIANTS", "")
rows = [(i["invariant"], i["necessary"], i["evidence"]) for i in INV["candidate_invariants_tested"]]
write_table(ws, 5, ["Invariant", "Necessary?", "Evidence"], rows, widths=[30, 18, 100])

# 238 Compiler Correctness
ws = new_sheet("238 RECFG-010 Correctness")
write_title(ws, "COMPILER CORRECTNESS", "")
r = block(ws, 5, "Status", "Compile(Theory)->UOC-IR->compiled representation, Decode(Compile(Theory))~=Theory: NOT ESTABLISHED for any nontrivial theory beyond UOC-IR's own schema (the trivial self-referential case). No domain theory (Maxwell, Einstein, etc.) was compiled through the full pipeline. NOT ATTEMPTED beyond the self-referential case.", height=100)

# 239 Compiler Completeness
ws = new_sheet("239 RECFG-011 Completeness")
write_title(ws, "COMPILER COMPLETENESS (19-point criterion)", "")
write_table(ws, 5, ["Criteria", "Verdict"], [
    ("1-4: kernel defined, non-uniqueness characterized, transformation defined, admissibility derived", "MET"),
    ("5: fixed-point/self-hosting closed", "PARTIALLY MET (narrow IR self-representability only)"),
    ("6-7: mathematical/spectral representation generated", "PARTIALLY MET"),
    ("8-16: geometry through observables", "NOT MET, each individually diagnosed"),
    ("17: no upstream depends on downstream observations", "MET, verified"),
    ("18: every node has a proof obligation", "MET"),
    ("19: every unresolved dependency explicitly identified", "MET"),
], widths=[80, 60])
r = 13
r = block(ws, r, "Verdict", "The compiler KERNEL is complete; the compiler as a whole is not, and is not claimed to be.", height=40)

# 240 Self-Hosting
ws = new_sheet("240 RECFG-012 Self-Hosting")
write_title(ws, "SELF-HOSTING", "")
r = block(ws, 5, "Result", "UOC-IR's own schema is representable as a UOC-IR instance (PROVEN, narrow). Full self-hosting Compile(UOC*)~=UOC* requires the entire pipeline to close, which it does not. OPEN, not falsified.", height=100)

# 241 Master MDCL
ws = new_sheet("241 RECFG-013 Master MDCL")
write_title(ws, "MASTER MDCL (DAG)", "")
rows = [(n["id"], n["object"], n["status"], ", ".join(n["parents"]), ", ".join(n["children"])) for n in MDCL["nodes"]]
write_table(ws, 5, ["ID", "Object", "Status", "Parents", "Children"], rows, widths=[30, 45, 40, 30, 30])

# 242 Master Equation Registry
ws = new_sheet("242 RECFG-014 Equation Reg")
write_title(ws, "MASTER EQUATION REGISTRY (minimized, partitioned)", "")
eq = load("uoc_master_equation_registry.json")
rows = []
for cat in ["primitive_definitions","compiler_rules","derived_equations","equivalent_representations","empirical_inputs"]:
    for item in eq[cat]:
        rows.append((cat, item))
write_table(ws, 5, ["Category", "Item"], rows, widths=[30, 130])

# 243 Master Proof Registry
ws = new_sheet("243 RECFG-015 Proof Reg")
write_title(ws, "MASTER PROOF / THEOREM REGISTRY", "")
rows = [(t["id"], t["statement"], t["status"]) for t in PROOF["theorems"]]
write_table(ws, 5, ["Theorem", "Statement", "Status"], rows, widths=[35, 90, 40])

# 244 Master Closure Matrix
ws = new_sheet("244 RECFG-016 Closure Matrix")
write_title(ws, "MASTER CLOSURE MATRIX", "")
rows = [(row["id"], row["status"], row["compiler_pass"], row["next_dependency"]) for row in CLOSURE["rows"]]
write_table(ws, 5, ["ID", "Status", "Compiler pass", "Next dependency"], rows, widths=[45, 40, 18, 90])

# 245 Historical Reclassification
ws = new_sheet("245 RECFG-017 Hist Reclass")
write_title(ws, "HISTORICAL RECLASSIFICATION", "Includes the Section 18 source-conflict finding")
r = block(ws, 5, "Critical finding -- SOURCE CONFLICT", HIST["critical_finding_SOURCE_CONFLICT"]["what_the_source_corpus_actually_shows"], height=140)
rows = [(t["result"], t["status"], str(t["still_valid"])) for t in HIST["reclassification_table"]]
write_table(ws, r, ["Result", "Status", "Still valid?"], rows, widths=[55, 55, 20])

# 246 Falsification Ledger
ws = new_sheet("246 RECFG-018 Falsification")
write_title(ws, "FALSIFICATION LEDGER", "")
rows = [(f["claim"], f["result"]) for f in FALS["falsified_this_run"]]
write_table(ws, 5, ["Claim", "Result"], rows, widths=[70, 90])

# 247 Final Closure
ws = new_sheet("247 RECFG-019 Final Closure")
write_title(ws, "FINAL CLOSURE STATUS", "")
ws["A5"] = ("Final Theorem (Universal Organizational Compiler Closure): OPEN / PROVEN NON-IDENTIFIABLE. "
            "The theorem presupposes a unique minimal kernel K* -- false (|F_N^derived_v2|=1,4,23, "
            "non-unique, growing). Not falsified either: a family-valued reformulation remains fully "
            "consistent with everything proven.\n\n"
            "Kernel complete and correctly reconfigured (TH-ARBS-001A/B correctly re-attached to R and "
            "N_orient(R) respectively; the Symmetric-cap-Acyclic incompatibility dissolved). Full "
            "compiler OPEN, blocked on three explicitly identified dependencies: lambda_gap "
            "(persistence), N-scaling/continuum behavior of F_N^derived_v2 (geometry), and the "
            "empirically-rejected gauge-automorphism route (gauge sector remains an input, not an "
            "output). None invented, none forced closed.")
ws["A5"].font = BODY_FONT
ws["A5"].alignment = WRAP
ws.merge_cells("A5:F5")
ws.row_dimensions[5].height = 220

wb.save(OUT)
print("saved", OUT)
print("total sheets:", len(wb.sheetnames))
