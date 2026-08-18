"""
Build v2.3 from v2.2: preserves all 197 existing sheets, appends the
UOC-C0-MINIMAL-SEED-CLOSURE-001 sheets (SEED-001 .. SEED-019).
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import json

SRC = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v2.2.xlsx"
OUT = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v2.3.xlsx"
RECON = "/home/user/URSP/reconstruction"

def load(name):
    with open(f"{RECON}/{name}") as f:
        return json.load(f)

CAND = load("seed_candidate_registry.json")
FPC = load("seed_fixed_point_catalog.json")
ISO = load("seed_isomorphism_classes.json")
CPLX = load("seed_complexity_comparison.json")
MIN = load("seed_minimality_audit.json")
SC = load("seed_self_compilation.json")
STAB = load("seed_stability.json")
UNIV = load("seed_universality.json")
FALS = load("seed_falsification.json")
DELTA = load("seed_delta_endaut.json")
RELISO = load("seed_relation_iso.json")

wb = openpyxl.load_workbook(SRC, data_only=False)

TITLE_FONT = Font(name="Arial", size=12, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def new_sheet(name):
    return wb.create_sheet(name)

def write_title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = WRAP

def write_table(ws, start_row, headers, rows, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + 1 + len(rows)

def block(ws, r, title, text, height=60):
    ws.cell(row=r, column=1, value=title).font = Font(bold=True)
    r += 1
    ws.cell(row=r, column=1, value=text)
    ws.cell(row=r, column=1).font = BODY_FONT
    ws.cell(row=r, column=1).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = height
    return r + 2


# SEED-001 Candidate Definitions
ws = new_sheet("SEED-001 Candidate Defs")
write_title(ws, "CANDIDATE DEFINITIONS", "Formal definitions and per-candidate findings, see seed_candidate_registry.json")
r = block(ws, 5, "Target independence", CAND["target_independence_declaration"], height=60)
rows = []
for name, c in CAND["candidates"].items():
    rows.append((name, c.get("definition",""), c.get("finding", c.get("identification_with_C",""))))
write_table(ws, r, ["Candidate", "Definition", "Finding"], rows, widths=[24, 60, 90])

# SEED-002 Distinction Enumeration
ws = new_sheet("SEED-002 Distinction Enum")
write_title(ws, "DISTINCTION (Delta) ENUMERATION -- End/Aut, N=1..6", "Formula validated by brute force for N<=4 (100% match).")
rows = [(d["N"], str(d["block_sizes"]), d["End_size_formula"], d["Aut_size_formula"],
         d.get("formula_matches_bruteforce", "n/a (N>4, formula only)")) for d in DELTA]
write_table(ws, 5, ["N", "Block sizes", "|End(Delta)|", "|Aut(Delta)|", "Formula==bruteforce"], rows, widths=[6,25,18,18,22])
r = 5 + len(rows) + 3
r = block(ws, r, "THEOREM-SEED-DELTA-001",
          "Aut(Delta) > 1 for EVERY partition shape tested at N>=2 (all integer partitions, N=1..6) "
          "-- bare distinction cannot canonically select a nonidentity transformation for N>=2. "
          "Sufficient only in the degenerate N<=1 case.", height=70)

# SEED-003 Relation Enumeration
ws = new_sheet("SEED-003 Relation Enum")
write_title(ws, "RELATION (R) ENUMERATION -- isomorphism classes, N=1..4 (exhaustive)", "")
rows = []
for N, d in RELISO.items():
    rows.append((d["N"], d["total_relations_raw"], d["total_isomorphism_classes"],
                 d["orbit_stabilizer_identity_verified_for_all_classes"],
                 d["functional_relations_raw_count"], d["rigid_isomorphism_classes_Aut_eq_1"],
                 d["rigid_functional_isomorphism_classes"]))
write_table(ws, 5, ["N", "Total relations (raw)", "Iso-classes", "Orbit-stabilizer verified",
                     "Functional (raw)", "Rigid classes (Aut=1)", "Rigid+functional"], rows, widths=[6,16,12,20,14,18,16])
r = 5 + len(rows) + 3
r = block(ws, r, "THEOREM-SEED-RELATION-001",
          "Rigid (Aut=1) functional relations exist for EVERY N>=1: the predecessor-chain construction "
          "tau(1)=1, tau(i)=i-1 is rigid, verified by brute-force automorphism check for N=1..7 and "
          "proven by hand for all N. Unlike Delta, R can be rigid -- the symmetry obstruction that "
          "blocks Delta does not block R.", height=80)

# SEED-004 Relation-Derived Operators
ws = new_sheet("SEED-004 Derived Operators")
write_title(ws, "D_R, T_R, C_R -- DERIVATION FROM R", "")
write_table(ws, 5, ["Operator", "Definition", "Domain of validity"], [
    ("D_R", "Coarsest equitable (1-WL/color-refinement stable) partition of X under R", "Every relation R (unconditional)"),
    ("T_R", "tau(x) = the unique y with R(x,y)", "R functional only (exactly N^N of 2^(N^2) relations)"),
    ("T_R (weaker)", "Aut(R): bijective automorphisms", "Every R, but generally non-unique (|Aut(R)|>1)"),
    ("C_R", "C_R(T) := 1[T(R)=R]", "Every R; tested against the weaker T(R) subset R (End-type) alternative"),
], widths=[16, 70, 60])

# SEED-005 Constraint Recovery
ws = new_sheet("SEED-005 Constraint Recovery")
write_title(ws, "CONSTRAINT RECOVERY", "")
r = block(ws, 5, "THM-SEED-CONSTRAINT-001",
          "{T : T(R)=R} is a strict subset of {T : T(R) subset R} in general (equality is strictly more "
          "restrictive than containment) -- a logical/degenerate proof, not requiring computation. This "
          "run computed the STRICT (Aut-type, bijective, T(R)=R) version exhaustively for all relations "
          "N=1..4. The weaker End-type monoid (T(R) subset R, not necessarily bijective) was NOT "
          "separately enumerated for pure relations this run -- recorded as an explicit scope gap.", height=100)

# SEED-006 Generative Operator
ws = new_sheet("SEED-006 Generative Operator")
write_title(ws, "Gamma_R = C_R o T_R o D_R -- THE COMPRESSION TEST", "")
r = block(ws, 5, "THM-SEED-COMPRESSION-001 (PROVEN)",
          "Gamma_R(R) is computable directly from R with ZERO additional stored information: D_R, T_R, "
          "C_R are all pure functions of R's own adjacency data (color-refinement classes, "
          "functional-successor-or-Aut(R), R-preservation predicate respectively). Verified by the "
          "complexity calculation in SEED-010: K(D_R,T_R,C_R | R) = 0 extra bits.", height=90)

# SEED-007 Fixed-Point Search
ws = new_sheet("SEED-007 Fixed-Point Search")
write_title(ws, "FIXED-POINT SEARCH -- Gamma_1 vs Gamma_infty", "")
write_table(ws, 5, ["N", "Total relations", "Iso-classes", "Trivial fixed", "Nontrivial fixed"], [
    (c["N"], c["total_relations_raw"], c["total_isomorphism_classes"], c["trivial_fixed_points"], c["nontrivial_fixed_points"])
    for c in FPC["catalog"]
], widths=[6,16,12,14,18])
r = 5 + len(FPC["catalog"]) + 3
r = block(ws, r, "Enumeration boundary", FPC["enumeration_boundary"]["reason"], height=90)
r = block(ws, r, "Gamma_1 vs Gamma_infty (idempotence discovery)", FPC["gamma_operator_note"], height=140)

# SEED-008 Isomorphism Classes
ws = new_sheet("SEED-008 Isomorphism Classes")
write_title(ws, "ISOMORPHISM CLASSES", "")
r = block(ws, 5, "Delta candidate (=integer partitions of N)", json.dumps(ISO["delta_candidate_isomorphism_classes"]["counts_by_N"]), height=40)
r = block(ws, r, "Relation candidate (all relations)", json.dumps(ISO["relation_candidate_isomorphism_classes"]["counts_by_N"]), height=40)
r = block(ws, r, "Gamma-fixed-point iso-classes", json.dumps(ISO["gamma_fixed_point_isomorphism_classes"]["counts_by_N"]), height=40)
r = block(ws, r, "Canonicality result", ISO["gamma_fixed_point_isomorphism_classes"]["canonicality_result"], height=110)

# SEED-009 Canonicality
ws = new_sheet("SEED-009 Canonicality")
write_title(ws, "CANONICALITY -- OUTCOME G", "")
r = block(ws, 5, "Conclusion", ISO["conclusion"], height=90)

# SEED-010 Complexity
ws = new_sheet("SEED-010 Complexity")
write_title(ws, "COMPLEXITY COMPARISON (N=4, declared encoding S1)", "")
nc = CPLX["numeric_comparison_at_N_equals_4"]
write_table(ws, 5, ["Quantity", "Bits"], [
    ("K(R)", nc["K_R"]), ("K(Delta)", nc["K_Delta"]), ("K(Delta,R) independent", nc["K_Delta_R_independent"]),
    ("K(D,T,C) independent, worst case", nc["K_D_T_C_independent_worst_case"]),
    ("K(D_R,T_R,C_R) derived, effective", nc["K_D_R_T_R_C_R_derived_effective"]),
], widths=[45, 15])
r = 12
r = block(ws, r, "Encoding scheme S1", CPLX["declared_encoding_family"]["name"], height=30)
r = block(ws, r, "Syntactic vs semantic compression", CPLX["syntactic_vs_semantic_compression"]["result"], height=90)

# SEED-011 Minimality
ws = new_sheet("SEED-011 Minimality")
write_title(ws, "MINIMALITY AUDIT (existence != uniqueness != minimality)", "")
r = block(ws, 5, "Existence", MIN["existence"]["status"], height=50)
r = block(ws, r, "Uniqueness", MIN["uniqueness"]["status"], height=60)
r = block(ws, r, "Minimality", MIN["minimality"]["status"], height=70)
r = block(ws, r, "Conflation check", MIN["conflation_check"], height=70)

# SEED-012 Self-Compilation
ws = new_sheet("SEED-012 Self-Compilation")
write_title(ws, "SELF-COMPILATION -- Compile(R*) ~ R*", "")
r = block(ws, 5, "Compile definition used", SC["compile_definition_used"], height=60)
r = block(ws, r, "Dynamical vs compiler fixed point", SC["dynamical_vs_compiler_fixed_point"]["relationship_found"], height=90)
r = block(ws, r, "Important negative finding (Gamma_1 not idempotent)", SC["important_negative_finding"]["result"], height=140)

# SEED-013 Stability
ws = new_sheet("SEED-013 Stability")
write_title(ws, "STABILITY (Hamming-1 perturbation, N=4, exhaustive)", "")
write_table(ws, 5, ["Metric", "Value"], [
    ("Isomorphism-class fixed points tested", STAB["results"]["n_isomorphism_class_fixed_points_tested"]),
    ("Total perturbation trials", STAB["results"]["total_perturbation_trials"]),
    ("Average fraction remaining fixed under 1 flip", STAB["results"]["average_fraction_of_1_flips_remaining_fixed"]),
    ("Fixed points destabilized by every flip", STAB["results"]["n_fixed_points_where_every_flip_destabilizes"]),
    ("Fixed points stable under every flip", STAB["results"]["n_fixed_points_locally_maximally_stable_survive_every_flip"]),
], widths=[55, 20])
r = 12
r = block(ws, r, "Asymptotic stability", STAB["classification"]["asymptotically_stable"], height=70)

# SEED-014 Universality
ws = new_sheet("SEED-014 Universality")
write_title(ws, "UNIVERSALITY", "")
r = block(ws, 5, "Abstract/discrete systems", "Functional relations = finite deterministic dynamical systems / single-symbol automata (direct identification). Multi-symbol automata = Sigma-indexed family of relations, no new primitive needed.", height=70)
r = block(ws, r, "Physical systems NOT tested", UNIV["physical_target_systems_NOT_tested"]["reason_not_tested"], height=140)
r = block(ws, r, "Conclusion", UNIV["universality_conclusion"], height=60)

# SEED-015 Falsification
ws = new_sheet("SEED-015 Falsification")
write_title(ws, "FALSIFICATION LEDGER", "")
rows = [(t["removed"], t["does_closure_survive"], t["conclusion"]) for t in FALS["structural_elimination"]]
write_table(ws, 5, ["Removed", "Closure survives?", "Conclusion"], rows, widths=[40, 70, 50])
r = 5 + len(rows) + 3
r = block(ws, r, "Type I (self-loop) vs Type III/IV (self-generation)", FALS["critical_distinction_self_relation_vs_self_generation_section_24"]["result"], height=120)

# SEED-016 Theorem Registry
ws = new_sheet("SEED-016 Theorem Registry")
write_title(ws, "THEOREM REGISTRY", "")
write_table(ws, 5, ["Theorem", "Status"], [
    ("THM-SEED-DELTA-001", "PROVEN -- Aut(Delta)>1 for N>=2 (all shapes N=1..6); insufficient except degenerate N<=1"),
    ("THM-SEED-RELATION-001", "PROVEN -- rigid functional relations exist for every N>=1 (general construction + brute force N=1..7)"),
    ("THM-SEED-DERIVATION-001", "D_R: DERIVED unconditionally. T_R: CONDITIONALLY DERIVED (functional case only)"),
    ("THM-SEED-CONSTRAINT-001", "PROVEN (logical) T(R)=R subset T(R) subset R; weaker End-type version not separately enumerated"),
    ("THM-SEED-COMPRESSION-001", "PROVEN -- Gamma_R(R) computable from R with zero extra stored information"),
    ("THM-SEED-FIXEDPOINT-001", "Existence PROVEN (N>=2). Gamma_1 idempotence FALSIFIED. Gamma_infty idempotence PROVEN"),
    ("THM-SEED-CANONICALITY-001", "PROVEN NON-UNIQUE -- 6/70/2462 nonisomorphic fixed points at N=2/3/4"),
    ("THM-SEED-MINIMALITY-001", "CONDITIONALLY DERIVED -- open under adversarial encoding; K(D_R,T_R,C_R)=K(R) PROVEN"),
    ("THM-SEED-SELF-COMPILATION-001", "PROVEN, conditional on Compile:=Gamma_infty"),
    ("THM-SEED-STABILITY-001", "CALCULATED -- 91.96% avg / 45.7% maximal local stability; asymptotic OPEN"),
    ("THM-SEED-UNIVERSALITY-001", "CONDITIONALLY DERIVED -- abstract PROVEN; physical OPEN by design"),
], widths=[32, 110])

# SEED-017 Closure Matrix
ws = new_sheet("SEED-017 Closure Matrix")
write_title(ws, "CLOSURE MATRIX", "")
write_table(ws, 5, ["Object", "Status"], [
    ("Bare distinction Delta as sufficient seed", "FALSIFIED for N>=2 -- PROVEN"),
    ("Pure relation R as sufficient seed", "SUFFICIENT -- existence PROVEN, N>=2"),
    ("Canonical/unique R*", "PROVEN NON-UNIQUE -- OUTCOME G"),
    ("Typed relation necessity", "NOT SHOWN NECESSARY -- reduces to a tuple of R"),
    ("Hypergraph necessity", "NOT SHOWN NECESSARY -- R already closes"),
    ("(D,T,C) irreducibility", "NOT SHOWN -- reduces to R under canonical/derived reading"),
    ("Gamma_1 (one-shot) idempotence", "FALSIFIED"),
    ("Gamma_infty (iterated) idempotence", "PROVEN"),
    ("Continuum limit / physical system tests", "NOT ATTEMPTED -- gated by unresolved canonicality, per stop condition"),
    ("Runs 1-12 (ARBS/spectral)", "PRESERVED, UNCHANGED, UNTOUCHED"),
], widths=[45, 95])

# SEED-018 MDCL
ws = new_sheet("SEED-018 MDCL")
write_title(ws, "MDCL (this run's layer)", "")
ws["A5"] = ("(D,T,C) [NOT SHOWN IRREDUCIBLE] < Gamma = C o T o D\n\n"
            "R subset X x X [SUFFICIENT, PROVEN] < D_R,T_R,C_R [DERIVED FROM R, 0 extra bits]\n"
            "  < Gamma_R [PROVEN COMPRESSIBLE] < R* = Gamma(R*) [EXISTS, PROVEN; NON-UNIQUE, PROVEN]\n\n"
            "Delta = (X,~) [PROVEN INSUFFICIENT for N>=2] -- separate, dead-end branch, not discarded, "
            "kept for the record per the preservation mandate.\n\n"
            "No proven map connects this seed layer to ARBS/G_k (runs 1-12) -- deliberately not attempted, "
            "per stop condition; canonicality is unresolved at this layer, so any such map would inherit "
            "an arbitrary choice among 2462+ nonisomorphic N=4 candidates.")
ws["A5"].font = BODY_FONT
ws["A5"].alignment = WRAP
ws.merge_cells("A5:F5")
ws.row_dimensions[5].height = 200

# SEED-019 Next Dependency
ws = new_sheet("SEED-019 Next Dependency")
write_title(ws, "EXACTLY ONE NEXT UNRESOLVED UPSTREAM DEPENDENCY", "")
ws["A4"] = ("Find a source-derived (not invented) selection principle that picks a unique member of "
            "the Gamma-fixed-point set -- or prove, for general N (this run reached only N<=4 "
            "exhaustively), that no such principle can exist and the seed is inherently a FAMILY, not "
            "an object. Until this resolves, no continuum limit, no physical-system test, and no "
            "downstream ARBS/spectral connection may be attempted from this seed layer without silently "
            "picking one of the many nonisomorphic candidates -- exactly the hidden-assumption failure "
            "this run's own stop condition exists to prevent.")
ws["A4"].font = Font(name="Arial", size=12, bold=True)
ws["A4"].alignment = WRAP
ws.merge_cells("A4:F4")
ws.row_dimensions[4].height = 200

wb.save(OUT)
print("saved", OUT)
print("total sheets:", len(wb.sheetnames))
