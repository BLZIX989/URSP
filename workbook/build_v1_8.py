"""
Build v1.8 from v1.7: preserves all 109 existing sheets, appends the C-004H
ARBS continuum compatibility recovery sheets.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import json

SRC = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v1.7.xlsx"
OUT = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v1.8.xlsx"
RECON = "/home/user/URSP/reconstruction"

def load(name):
    with open(f"{RECON}/{name}") as f:
        return json.load(f)

REG = load("C004H_registry.json")
DIAM = load("c004h_diameter.json")
INV = load("c004h_metric_invariance.json")
REF = load("c004f_refinement.json")
PUSH = load("c004f_pushforward.json")
DIR_ = load("c004f_dirichlet.json")

wb = openpyxl.load_workbook(SRC, data_only=False)

TITLE_FONT = Font(name="Arial", size=12, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def new_sheet(name):
    return wb.create_sheet(name)


def write_title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = WRAP


def write_table(ws, start_row, headers, rows, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + 1 + len(rows)


# 110 Overview
ws = new_sheet("110 C-004H Overview")
write_title(ws, "C-004H -- ARBS CONTINUUM COMPATIBILITY RECOVERY",
            "Governing question: " + REG["governing_question"])
write_table(ws, 5, ["New finding this run"], [[x] for x in REG["source_exhaustion_search"]["new_findings_this_run"]], widths=[140])

# 111 Source Search
ws = new_sheet("111 C-004H Source Search")
write_title(ws, "SOURCE EXHAUSTION SEARCH", REG["source_exhaustion_search"]["scope"])
r = 5
for i, finding in enumerate(REG["source_exhaustion_search"]["new_findings_this_run"], 1):
    ws.cell(row=r, column=1, value=f"Finding {i}").font = Font(bold=True)
    r += 1
    ws.cell(row=r, column=1, value=finding).font = BODY_FONT
    ws.cell(row=r, column=1).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 60
    r += 2

# 112 Refinement structure (carried forward)
ws = new_sheet("112 C-004H Refinement")
write_title(ws, "C-004H.1 -- ARBS REFINEMENT STRUCTURE", "Carried forward from C-004F.5, unaffected by this run.")
d = REG["C004H_1_refinement_structure"]
write_table(ws, 5, ["Field", "Value"], [
    ("Result", d["result"]), ("Classification", d["classification"]),
    ("Every edge has R(e)?", d["does_every_edge_have_R(e)"]), ("Status", d["status"]),
], widths=[26, 110])

# 113 Edge Lengths
ws = new_sheet("113 C-004H Edge Lengths")
write_title(ws, "C-004H.2 -- EDGE LENGTH RECOVERY", "")
d = REG["C004H_2_edge_length_recovery"]
ws["A5"] = d["search_result"]
ws["A5"].font = BODY_FONT; ws["A5"].alignment = WRAP
ws.merge_cells("A5:F5"); ws.row_dimensions[5].height = 70
ws["A7"] = f"Status: {d['status']}"
ws["A7"].font = Font(bold=True)

# 114 Edge Scaling Test
ws = new_sheet("114 C-004H Edge Scaling")
write_title(ws, "C-004H.3 -- EDGE SCALING TEST", "Carried forward from C-004G, unaffected.")
d = REG["C004H_3_edge_scaling_test"]
ws["A5"] = d["result"]
ws["A5"].font = BODY_FONT; ws["A5"].alignment = WRAP
ws.merge_cells("A5:F5"); ws.row_dimensions[5].height = 50
ws["A7"] = f"Status: {d['status']}"
ws["A7"].font = Font(bold=True, color="C00000")

# 115 Metric Candidates / Diameter (the centerpiece new result)
ws = new_sheet("115 C-004H Metric Diameter")
write_title(ws, "C-004H.4-6 -- METRIC RECOVERY, INVARIANCE, DIAMETER (CENTERPIECE FINDING)",
            "Candidate: " + REG["C004H_4_5_6_metric_recovery_invariance_diameter"]["candidate"])
rows = [(k, v["N"], v["diffusion_diam"], v["resistance_diam"], v["identity_check_max|d^2-Reff|"])
        for k, v in DIAM.items()]
r = write_table(ws, 5, ["Shell", "N", "Diffusion diameter", "R_eff diameter", "|d^2-R_eff| residual"], rows, widths=[8, 8, 20, 18, 20])
r += 2
ws.cell(row=r, column=1, value="Diameter verdict").font = Font(bold=True)
r += 1
ws.cell(row=r, column=1, value=REG["C004H_4_5_6_metric_recovery_invariance_diameter"]["diameter_stability_verdict"])
ws.cell(row=r, column=1).font = BODY_FONT; ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); ws.row_dimensions[r].height = 40
r += 2
ws.cell(row=r, column=1, value="Induced metric invariance -- violation trend (old-vertex pairs)").font = Font(bold=True)
r += 1
rows2 = [(k, v["max_|d_{k+1}(vi,vj)-d_k(vi,vj)|_over_old_pairs"], v["ratio_prev_over_curr"]) for k, v in INV.items()]
r = write_table(ws, r, ["Transition", "Max violation", "Ratio prev/curr"], rows2, widths=[16, 20, 20])
r += 2
ws.cell(row=r, column=1, value="Invariance verdict").font = Font(bold=True)
r += 1
ws.cell(row=r, column=1, value=REG["C004H_4_5_6_metric_recovery_invariance_diameter"]["induced_metric_invariance_verdict"])
ws.cell(row=r, column=1).font = BODY_FONT; ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); ws.row_dimensions[r].height = 55
r += 3
ws.cell(row=r, column=1, value=REG["C004H_4_5_6_metric_recovery_invariance_diameter"]["caveat"])
ws.cell(row=r, column=1).font = NOTE_FONT; ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); ws.row_dimensions[r].height = 40
r += 2
ws.cell(row=r, column=1, value=f"Status: {REG['C004H_4_5_6_metric_recovery_invariance_diameter']['status']}").font = Font(bold=True, size=12, color="1F7A1F")

# 116 Measure Pushforward (carried forward + tension noted)
ws = new_sheet("116 C-004H Measure")
write_title(ws, "C-004H.7 -- MEASURE RECOVERY", "Carried forward from C-004F.6/7, unaffected by the metric findings.")
d = REG["C004H_7_measure_recovery"]
ws["A5"] = d["result"]
ws["A5"].font = BODY_FONT; ws["A5"].alignment = WRAP
ws.merge_cells("A5:F5"); ws.row_dimensions[5].height = 40
ws["A7"] = "Critical tension with the C-004H.4-6 metric finding:"
ws["A7"].font = Font(bold=True)
ws["A8"] = d["critical_tension_with_C004H_4_6"]
ws["A8"].font = BODY_FONT; ws["A8"].alignment = WRAP
ws.merge_cells("A8:F8"); ws.row_dimensions[8].height = 80
ws["A10"] = f"Status: {d['status']}"
ws["A10"].font = Font(bold=True, color="C00000")

# 117 Dirichlet Form (carried forward)
ws = new_sheet("117 C-004H Dirichlet")
write_title(ws, "C-004H.8 -- DIRICHLET FORM COMPATIBILITY", "Carried forward from C-004F.4, unaffected.")
d = REG["C004H_8_dirichlet_form_compatibility"]
ws["A5"] = d["result"]
ws["A5"].font = BODY_FONT; ws["A5"].alignment = WRAP
ws.merge_cells("A5:F5"); ws.row_dimensions[5].height = 40
ws["A7"] = f"Status: {d['status']}"
ws["A7"].font = Font(bold=True)

# 118 Continuum Compatibility Theorem
ws = new_sheet("118 C-004H Continuum Compat")
write_title(ws, "THEOREM C-004H.9 -- CONTINUUM COMPATIBILITY", REG["theorem_C004H_9"]["attempted_statement"])
r = 6
write_table(ws, r, ["Component", "Result"], list(REG["theorem_C004H_9"]["componentwise_result"].items()), widths=[36, 100])
r = r + len(REG["theorem_C004H_9"]["componentwise_result"]) + 3
ws.cell(row=r, column=1, value="Synthesis").font = Font(bold=True, size=12)
r += 1
ws.cell(row=r, column=1, value=REG["theorem_C004H_9"]["synthesis"])
ws.cell(row=r, column=1).font = BODY_FONT; ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); ws.row_dimensions[r].height = 160
r += 3
ws.cell(row=r, column=1, value=f"PROMOTED STATUS: {REG['theorem_C004H_9']['promoted_status']}")
ws.cell(row=r, column=1).font = Font(bold=True, size=12, color="C00000")
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); ws.row_dimensions[r].height = 60

# 119 Architectural Boundary
ws = new_sheet("119 C-004H Boundary")
write_title(ws, "ARCHITECTURAL BOUNDARY RECORDED", "")
write_table(ws, 5, ["Object", "Status"], list(REG["architectural_boundary_recorded"].items()), widths=[36, 100])

# 120 Theorem / Falsification / Closure
ws = new_sheet("120 C-004H Closure Ledger")
write_title(ws, "TARGET INDEPENDENCE, FALSIFICATION LEDGER, CLOSURE", "")
r = 5
ws.cell(row=r, column=1, value="Target independence").font = Font(bold=True)
r += 1
ws.cell(row=r, column=1, value=REG["target_independence_audit"]["result"])
ws.cell(row=r, column=1).font = BODY_FONT; ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); ws.row_dimensions[r].height = 40
r += 3
ws.cell(row=r, column=1, value="Falsification ledger addition").font = Font(bold=True)
r += 1
ws.cell(row=r, column=1, value=REG["falsification_ledger_additions"]["unit_edge_hop_metric_as_THE_canonical_ARBS_metric"])
ws.cell(row=r, column=1).font = BODY_FONT; ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); ws.row_dimensions[r].height = 55

# 121 Next Dependency
ws = new_sheet("121 C-004H Next Dependency")
write_title(ws, "EXACTLY ONE NEXT UNRESOLVED UPSTREAM DEPENDENCY", "")
ws["A4"] = REG["next_dependency"]
ws["A4"].font = Font(name="Arial", size=12, bold=True)
ws["A4"].alignment = WRAP
ws.merge_cells("A4:F4")
ws.row_dimensions[4].height = 180

wb.save(OUT)
print("saved", OUT)
print("total sheets:", len(wb.sheetnames))
