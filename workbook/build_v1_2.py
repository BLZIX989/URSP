"""
Build v1.2 from v1.1: preserves all 56 existing sheets, appends the
ARBS-CONSTRUCTION-001 graph reconstruction results and the re-executed
C-004B / OPEN-020E.2 sheets now that complete spectra exist.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import json
import numpy as np

SRC = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v1.1.xlsx"
OUT = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v1.2.xlsx"
RECON = "/home/user/URSP/reconstruction"

with open(f"{RECON}/full_analysis_results.json") as f:
    FULL = json.load(f)
with open(f"{RECON}/proof_obligations_P1_P11.json") as f:
    PO = json.load(f)
with open(f"{RECON}/checkpoint_comparison.json") as f:
    CKPT = json.load(f)
with open(f"{RECON}/reconstruction_audit.json") as f:
    AUDIT = json.load(f)

wb = openpyxl.load_workbook(SRC, data_only=False)

TITLE_FONT = Font(name="Arial", size=12, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def new_sheet(name):
    return wb.create_sheet(name)


def write_title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = WRAP


def write_table(ws, start_row, headers, rows, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + 1 + len(rows)


# ---------------------------------------------------------------------
# 57 ARBS Graph Construction
# ---------------------------------------------------------------------
ws = new_sheet("57 ARBS Graph Construction")
write_title(ws, "ARBS-CONSTRUCTION-001 -- SOURCE-SPECIFIED GRAPH CONSTRUCTION",
            "S_k = L_k sqcup R_k, |L_k|=|R_k|=2^k. Intra-shell: G[S_k]=K_{2^k,2^k}. "
            "Inter-shell: E(S_{k-1},S_k)=R_{k-1} x L_k. G_n = union_{k=0}^n S_k. "
            "Provenance note: this exact rule was supplied verbatim as directive text in this "
            "task; an independent search of the newly supplied 9-workbook + whitepaper corpus "
            "found no literal occurrence of it under ARBS-DEF-001/ARBS-SCALE-001-003 (those define "
            "ARBS as an Object/Operator-node bipartite typed-dependency graph, a different object). "
            "The corpus's own internal 'Audit and Recommended Revisions' document states the "
            "shell-hierarchy replacement rule f(S_m) 'is asserted but never specified anywhere in "
            "the eight-document corpus (confirmed by exhaustive keyword and functional search)' and "
            "that the source's own tested selector 'converges to near-complete graph, not the "
            "claimed Recursive Bipartite Shell Graph.' Despite this documentation gap, the rule is "
            "fully specified in this run's directive and is executed exactly as given below; its "
            "numerical output is independently and decisively confirmed against 8 checkpoint values "
            "(see sheet 58) to double-precision accuracy, which is the actual evidence for its "
            "correctness, not the corpus citation.")
rows = []
for k, meta in json.load(open(f"{RECON}/artifact_metadata.json")).items():
    rows.append((k, meta["vertex_count"], meta["edge_count"], meta["status"]))
write_table(ws, 5, ["Shell", "N (vertices)", "E (edges)", "Status"], rows, widths=[10, 16, 12, 22])

# ---------------------------------------------------------------------
# 58 Spectral Reconstruction Audit (P1-P11 + checkpoints)
# ---------------------------------------------------------------------
ws = new_sheet("58 Spectral Recon Audit")
write_title(ws, "SPECTRAL RECONSTRUCTION AUDIT -- P1-P11 + CHECKPOINT COMPARISON",
            "Every check computed directly from the constructed graphs (build_arbs.py); "
            "PASS/FAIL only, no qualitative judgments.")
r = 5
ws.cell(row=r, column=1, value="Proof obligations P1-P11 (per shell)").font = Font(bold=True)
r += 1
po_headers = ["Shell"] + list(next(iter(PO.values())).keys())
po_rows = [[k] + [v[h] for h in po_headers[1:]] for k, v in PO.items()]
r = write_table(ws, r, po_headers, po_rows, widths=[8] + [16] * (len(po_headers) - 1))
r += 2

ws.cell(row=r, column=1, value="Numeric checkpoint comparison (checkpoint-derived vs. independently reconstructed)").font = Font(bold=True)
r += 1
r = write_table(ws, r,
    ["Checkpoint", "Reconstructed value", "v1.0 checkpoint value", "Abs. diff", "Match (<1e-6)?"],
    [
        ("G3 sigma16", CKPT["G3_sigma16_calculated"], CKPT["G3_sigma16_checkpoint"],
         abs(CKPT["G3_sigma16_calculated"] - CKPT["G3_sigma16_checkpoint"]), CKPT["G3_sigma16_match"]),
        ("G3 sigma17", CKPT["G3_sigma17_calculated"], CKPT["G3_sigma17_checkpoint"],
         abs(CKPT["G3_sigma17_calculated"] - CKPT["G3_sigma17_checkpoint"]), CKPT["G3_sigma17_match"]),
        ("G4 sigma16=sigma17 (degenerate)", CKPT["G4_sigma16_calculated"], CKPT["G4_checkpoint"],
         abs(CKPT["G4_sigma16_calculated"] - CKPT["G4_checkpoint"]), CKPT["G4_value_match"]),
    ] + [
        (f"{k} t_H", v["t_H_calculated"], v["t_H_checkpoint"], v["t_H_abs_diff"], v["t_H_abs_diff"] < 1e-6)
        for k, v in FULL["per_shell"].items()
    ] + [
        (f"{k} lambda_c", v["lambda_c_calculated"], v["lambda_c_checkpoint"], v["lambda_c_abs_diff"], v["lambda_c_abs_diff"] < 1e-6)
        for k, v in FULL["per_shell"].items()
    ],
    widths=[28, 22, 22, 16, 14])

r += 1
ws.cell(row=r, column=1,
        value="All 8 checkpoint classes match to double-precision tolerance (<2e-10 absolute, typically ~1e-13). "
              "t_H and lambda_c were solved independently from the complete spectrum via root-finding on "
              "p_A(t)=Tr(P_+ e^{-2tL})/rank(P_+)=1/2, NOT copied from the v1.0 checkpoint sheet.")
ws.cell(row=r, column=1).font = NOTE_FONT
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
ws.row_dimensions[r].height = 40

# ---------------------------------------------------------------------
# 59 Full Spectral Arrays (Laplacian + Dirac positive)
# ---------------------------------------------------------------------
ws = new_sheet("59 Full Spectral Arrays")
write_title(ws, "COMPLETE LAPLACIAN AND POSITIVE DIRAC SPECTRA, G0-G4",
            "Ascending order. Full matrices (A, B, L, D_G) are exported as CSV+NumPy under "
            "reconstruction/{adjacency,incidence,laplacian,dirac}/ in the repository (too large "
            "for convenient spreadsheet embedding at G4 scale: D_G is 573x573).")
col = 1
row0 = 5
for n in range(5):
    key = f"G{n}"
    lam = np.load(f"{RECON}/spectra/{key}_laplacian.npy")
    pos = np.load(f"{RECON}/spectra/{key}_dirac_positive.npy")
    ws.cell(row=row0, column=col, value=f"{key} Laplacian eigenvalues (ascending, incl. zero mode)").font = HEADER_FONT
    ws.cell(row=row0, column=col).fill = HEADER_FILL
    ws.cell(row=row0, column=col + 1, value=f"{key} positive Dirac singular values (ascending)").font = HEADER_FONT
    ws.cell(row=row0, column=col + 1).fill = HEADER_FILL
    for i, val in enumerate(lam):
        ws.cell(row=row0 + 1 + i, column=col, value=float(val))
    for i, val in enumerate(pos):
        ws.cell(row=row0 + 1 + i, column=col + 1, value=float(val))
    ws.column_dimensions[get_column_letter(col)].width = 26
    ws.column_dimensions[get_column_letter(col + 1)].width = 30
    col += 3

# ---------------------------------------------------------------------
# 60 C-004B Re-Execution
# ---------------------------------------------------------------------
ws = new_sheet("60 C-004B Re-Execution")
write_title(ws, "C-004B RE-EXECUTED WITH COMPLETE SPECTRA",
            "N_H,k = rank(P_H,k) = #{Laplacian eigenvalues >= lambda_c,k} -- now directly computable "
            "for every shell (previously OPEN in v1.1 for lack of lambda_max).")
rows = [(k, v["N"], v["delta_spec_lambda1"], v["lambda_max"], v["lambda_c_calculated"],
         v["lambda_max"] < v["lambda_c_calculated"], v["N_H_rank_P_H"])
        for k, v in FULL["per_shell"].items()]
r = write_table(ws, 5,
    ["Shell", "N", "delta_spec = lambda_1(L)", "lambda_max(L)", "lambda_c=1/t_H", "lambda_max < lambda_c ?", "N_H = rank(P_H)"],
    rows, widths=[8, 8, 22, 16, 18, 20, 16])
r += 1
ws.cell(row=r, column=1,
        value="RESULT: lambda_max(L_k) < lambda_c,k for every shell G0-G4 -> N_H(G_k)=0 for all k=0..4. "
              "The canonical persistence threshold lambda_c=1/t_H does not merely fail to select the "
              "G3 32-boundary (v1.1 finding) -- it lies entirely above the top of the spectrum for "
              "every tested shell, producing an EMPTY horizon sector everywhere. "
              "C-004B = FALSIFIED FOR 32-BOUNDARY COMPATIBILITY is confirmed and strengthened.")
ws.cell(row=r, column=1).font = Font(bold=True)
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
ws.row_dimensions[r].height = 55

r += 2
ws.cell(row=r, column=1, value="G3 boundary rank tests").font = Font(bold=True)
r += 1
gb = FULL["G3_boundary"]
r = write_table(ws, r,
    ["Object", "Value"],
    [
        ("lambda16(G3)", gb["G3_lambda16"]),
        ("lambda17(G3)", gb["G3_lambda17"]),
        ("rank(1_[lambda16,lambda17)(L_3))", gb["rank_P_boundary_1_[lambda16,lambda17)_on_Laplacian"]),
        ("sigma16(G3)", gb["G3_sigma16"]),
        ("sigma17(G3)", gb["G3_sigma17"]),
        ("rank of projector onto first 16 nonzero +-sigma Dirac pairs", gb["rank_of_first_16_nonzero_Dirac_pairs_(+-sigma_1..sigma_16)"]),
    ], widths=[50, 24])
r += 1
ws.cell(row=r, column=1, value=gb["interpretation"])
ws.cell(row=r, column=1).font = NOTE_FONT
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
ws.row_dimensions[r].height = 90

# ---------------------------------------------------------------------
# 61 OPEN-020E.2 Re-Audit
# ---------------------------------------------------------------------
ws = new_sheet("61 OPEN-020E.2 Re-Audit")
write_title(ws, "OPEN-020E.2 RE-AUDIT WITH SOURCE-CONFIRMED Pi_0 DEFINITION",
            "Corpus search (5ec46f55 workbook, sheet 54_SOURCE_TEXT_CORPUS) recovered the genuine "
            "source definition: Pi_0 = (delta_spec - lambda_c) / sigma, with delta_spec=lambda_1(L) "
            "(CERTIFIED, B-002) and sigma = entropy production (CERTIFIED as existing and >=0 by an "
            "H-theorem, B-006), but sigma is defined dynamically as sigma=-dF/dt along a solution "
            "P(t) of the Fokker-Planck-type evolution -- NOT reduced to a static per-graph scalar "
            "anywhere in the source material. No initial condition P(0) or evaluation time is fixed "
            "by any registry entry found.")
r = write_table(ws, 6,
    ["Required object", "Status", "Detail"],
    [
        ("delta_spec(G_k) = lambda_1(L_k)", "NOW COMPUTABLE",
         "Directly available from the complete spectra reconstructed this run (sheet 59). Values: "
         + "; ".join(f"{k}={v['delta_spec_lambda1']:.6f}" for k, v in FULL["per_shell"].items())),
        ("lambda_c,k", "COMPUTABLE (already available)", "Canonical mapping lambda_c=1/t_H, sheet 60. "
         "Note: source Bridge B-007 flags the RIGOROUS derivation of lambda_c (C-004, distinct from "
         "the 32-boundary-selection question of C-004B) as itself OPEN -- the 1/t_H mapping is used "
         "as the project's working convention, not as a source-certified closed form."),
        ("sigma_k (entropy production, per shell, static number)", "STILL UNDEFINED",
         "Source defines sigma = -dF/dt dynamically (H-theorem, B-006, CERTIFIED as a theorem about "
         "existence and sign) but gives no per-graph static reduction: no fixed initial distribution "
         "P(0), no fixed evaluation time, and no closed-form spectral expression is recorded anywhere "
         "in the 9-workbook + whitepaper corpus. Computing a specific number here would require "
         "choosing P(0) and t -- exactly the invention this run's governing instructions forbid."),
    ], widths=[36, 26, 70])
r += 2
ws.cell(row=r, column=1, value="OPEN-020E.2 OUTCOME (unchanged: D)").font = Font(bold=True, color="C00000")
r += 1
ws.cell(row=r, column=1,
        value="D. Pi_0(G_k) still cannot be evaluated for any k. The blocking object has narrowed from "
              "'delta_spec and sigma_k both undefined' (v1.1) to 'sigma_k (entropy production) alone "
              "undefined as a static per-shell number' -- delta_spec is now fully computed. "
              "argmax_k Pi_0(G_k)=3 remains untested; outcome A is not forced.")
ws.cell(row=r, column=1).font = BODY_FONT
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
ws.row_dimensions[r].height = 55

# ---------------------------------------------------------------------
# 62 Status Update (this run)
# ---------------------------------------------------------------------
ws = new_sheet("62 Status Update (Run 2)")
write_title(ws, "STATUS UPDATE -- ARBS-CONSTRUCTION-001", "")
r = write_table(ws, 4,
    ["Object", "v1.1 status", "v1.2 status (this run)", "Basis"],
    [
        ("ARBS-FULL-SPEC-001 / graph construction", "checkpoint / internally consistent (no graph existed)",
         "DERIVED / VERIFIED", "P1-P11 all PASS for G0-G4 (sheet 58); construction rule matches all "
         "structural checkpoints by algebraic necessity (Section 5/6) and the reconstructed positive "
         "Dirac spectrum matches all 8 independent numeric checkpoints to <2e-10 (sheet 58)."),
        ("Spectral reconstruction (complete {lambda_k,j})", "OPEN (missing)", "DERIVED / VERIFIED",
         "Complete Laplacian and positive-Dirac spectra now exist for G0-G4 (sheet 59)."),
        ("C-004B", "FALSIFIED FOR 32-BOUNDARY COMPATIBILITY (partial: lambda_max unknown)",
         "FALSIFIED FOR 32-BOUNDARY COMPATIBILITY (confirmed, strengthened: N_H(G_k)=0 for ALL k=0..4)",
         "sheet 60"),
        ("N_H(G0..G4)", "OPEN (missing lambda_max)", "CALCULATED: 0, 0, 0, 0, 0", "sheet 60"),
        ("rank(P_boundary), G3", "OPEN", "CALCULATED: 1 (Laplacian half-open interval); 32 (first-16 "
         "nonzero +-sigma Dirac pairs, the correct reading of the '32-state' language)", "sheet 60"),
        ("OPEN-020E.2 / Pi_0(G_k)", "Outcome D (delta_spec AND sigma_k both undefined)",
         "Outcome D (delta_spec now computed; sigma_k -- entropy production -- remains undefined "
         "as a static number)", "sheet 61"),
        ("G* (shell selection)", "UNRESOLVED", "STILL UNRESOLVED",
         "No computed quantity (N_H, delta_spec, or any other) uniquely distinguishes G3 from "
         "G0/G1/G2/G4 this run; N_H is identically 0 for every shell. No uniqueness theorem exists "
         "in source. Per Section 21: not promoted because a numerical result looks favorable."),
    ], widths=[30, 40, 46, 60])

wb.save(OUT)
print("saved", OUT)
print("total sheets:", len(wb.sheetnames))
