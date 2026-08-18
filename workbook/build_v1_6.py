"""
Build v1.6 from v1.5: preserves all 88 existing sheets, appends the C-004F
degree-weighted thermodynamic continuum recovery sheets.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import json

SRC = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v1.5.xlsx"
OUT = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v1.6.xlsx"
RECON = "/home/user/URSP/reconstruction"

def load(name):
    with open(f"{RECON}/{name}") as f:
        return json.load(f)

REG = load("C004F_registry.json")
DM = load("c004f_degree_measure.json")
GEN = load("c004f_generator.json")
DIR = load("c004f_dirichlet.json")
REF = load("c004f_refinement.json")
PUSH = load("c004f_pushforward.json")
ENT = load("c004f_entropy_consequence.json")

wb = openpyxl.load_workbook(SRC, data_only=False)

TITLE_FONT = Font(name="Arial", size=12, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def new_sheet(name):
    return wb.create_sheet(name)


def write_title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = WRAP


def write_table(ws, start_row, headers, rows, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + 1 + len(rows)


# 89 Overview
ws = new_sheet("89 C-004F Overview")
write_title(ws, "C-004F -- DEGREE-WEIGHTED THERMODYNAMIC CONTINUUM RECOVERY",
            "Recovery order executed exactly: F.1 degree measure -> F.2 weighted Hilbert space -> "
            "F.3 weighted generator -> F.4 Dirichlet form -> F.5 refinement map -> F.6 pushforward -> "
            "F.7 convergence test -> [F.8-F.10 blocked, not attempted].")
write_table(ws, 5, ["Step", "Status"], list(REG["recovery_order_status"].items()), widths=[36, 90])

# 90 Degree Measures
ws = new_sheet("90 C-004F Degree Measures")
write_title(ws, "DEGREE MEASURES pi_k(i)=d_i/(2|E_k|)", "G0-G9 (extended per Section 17).")
rows = [(k, v["N"], v["E"], v["pi_min"], v["pi_max"], v["entropy_H(pi)"], v["entropy_H(uniform)"],
         v["||pi-u||_1"], v["||pi-u||_2"]) for k, v in DM.items()]
write_table(ws, 4, ["Shell", "N", "E", "pi_min", "pi_max", "H(pi)", "H(uniform)", "||pi-u||_1", "||pi-u||_2"],
            rows, widths=[8, 8, 10, 12, 12, 14, 14, 14, 14])

# 91 Stationary Distributions
ws = new_sheet("91 C-004F Stationary Distrib")
write_title(ws, "STATIONARITY OF pi_k UNDER Q_k*", "")
rows = [(k, v["Qstar_pi_norm_(should_be_0)"], v["Qstar_d_norm_(should_be_0)"], v["conservation_residual_1T_Qstar_(should_be_0)"])
        for k, v in GEN.items()]
write_table(ws, 4, ["Shell", "||Q*pi|| (=0 expect)", "||Q*d|| (=0 expect)", "||1^T Q*|| (=0 expect)"], rows, widths=[8, 20, 20, 20])

# 92 Weighted Hilbert Spaces
ws = new_sheet("92 C-004F Weighted Hilbert")
write_title(ws, "WEIGHTED HILBERT SPACE H_k = L^2(V_k, mu_k)", "")
ws["A5"] = f"Source support: {REG['C004F_2_weighted_hilbert_space']['source_support']}"
ws["A5"].font = BODY_FONT; ws["A5"].alignment = WRAP
ws.merge_cells("A5:E5"); ws.row_dimensions[5].height = 60
ws["A7"] = f"Positivity/nondegeneracy: {REG['C004F_2_weighted_hilbert_space']['positivity_nondegeneracy']}"
ws["A7"].font = BODY_FONT; ws["A7"].alignment = WRAP
ws.merge_cells("A7:E7"); ws.row_dimensions[7].height = 40
ws["A9"] = f"Status: {REG['C004F_2_weighted_hilbert_space']['status']}"
ws["A9"].font = Font(bold=True)

# 93 Thermodynamic Generators
ws = new_sheet("93 C-004F Generators")
write_title(ws, "WEIGHTED GENERATOR Q_k, Q_k*", "Q_ij=A_ij/d_i (i!=j), Q_ii=-1; Q*=Q^T acts on column P.")
rows = [(k, v["weighted_inner_product_positive_definite"],
         v["T(t)_stochastic_checks"]["t=1.0"]["row_stochastic"],
         v["T(t)_stochastic_checks"]["t=1.0"]["nonneg_entries"],
         v["T(t)_stochastic_checks"]["t=1.0"]["max_row_sum_dev"])
        for k, v in GEN.items()]
write_table(ws, 4, ["Shell", "Weighted IP positive-def?", "exp(Q) row-stochastic (t=1)?", "nonneg entries?", "max row-sum dev"],
            rows, widths=[8, 22, 26, 16, 18])

# 94 Detailed Balance
ws = new_sheet("94 C-004F Detailed Balance")
write_title(ws, "DETAILED BALANCE VERIFICATION", "pi_i Q_ij = pi_j Q_ji")
rows = [(k, v["detailed_balance_residual_(should_be_0)"]) for k, v in GEN.items()]
write_table(ws, 4, ["Shell", "||pi_i Q_ij - pi_j Q_ji||_F (=0 expect)"], rows, widths=[8, 40])

# 95 Dirichlet Forms
ws = new_sheet("95 C-004F Dirichlet Forms")
write_title(ws, "DIRICHLET FORM E_k(f,g) = (1/(4|E_k|)) sum A_ij(f_i-f_j)(g_i-g_j) = <f,-Q_kg>_mu",
            REG["C004F_4_dirichlet_form"]["numeric_verification"])
rows = [(k, v["verified_equals_<f,-Qg>_mu_max_residual"], v["E_k(const,const)_(should_be_0)"])
        for k, v in DIR.items()]
write_table(ws, 5, ["Shell", "max |E_k(f,g)-<f,-Qg>_mu| residual", "E_k(1,1)"], rows, widths=[8, 34, 14])

# 96 Refinement Maps
ws = new_sheet("96 C-004F Refinement Maps")
write_title(ws, "ARBS REFINEMENT MAP rho_k: G_k -> G_k+1",
            REF[list(REF.keys())[0]]["map_type"] + " -- " + REF[list(REF.keys())[0]]["metric_rescaling_status"][:150] + "...")
rows = [(k, v["inclusion_holds_(G_k_is_induced_subgraph_of_G_k+1)"], v["num_vertices_with_changed_degree"],
         v["changed_vertices_are_exactly_R_k"], v["R_k_degree_increase_(all_equal_2^(n+1)?)"]["uniform"],
         v["R_k_degree_increase_(all_equal_2^(n+1)?)"]["expected_increase"])
        for k, v in REF.items()]
write_table(ws, 5, ["Transition", "Inclusion holds?", "#vertices degree-changed", "Changed = R_k exactly?", "Increase uniform?", "Increase amount"],
            rows, widths=[16, 14, 20, 20, 16, 14])
r = 5 + len(rows) + 2
ws.cell(row=r, column=1, value="Metric rescaling status (all transitions)").font = Font(bold=True)
r += 1
ws.cell(row=r, column=1, value=REG["C004F_5_refinement_map"]["metric_rescaling_status"])
ws.cell(row=r, column=1).font = BODY_FONT
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
ws.row_dimensions[r].height = 70

# 97 Measure Pushforward
ws = new_sheet("97 C-004F Pushforward")
write_title(ws, "MEASURE PUSHFORWARD (rho_k)_* pi_k VS pi_k+1",
            "Section 17 caveat: exact finite-shell (G0-G9) results, NOT a k->infinity proof.")
rows = [(k, v["total_variation_distance"], v["L1_distance"], v["L2_distance"], v["mass_fraction_pi_{k+1}_on_new_vertices"])
        for k, v in PUSH.items()]
write_table(ws, 4, ["Transition", "TV distance", "L1 distance", "L2 distance", "Mass fraction on new vertices"],
            rows, widths=[16, 14, 14, 14, 26])
r = 4 + len(rows) + 2
ws.cell(row=r, column=1,
        value="CRITICAL FINDING: TV distance stabilizes at a strictly POSITIVE constant (~0.625), converging "
              "geometrically (each deviation ~4x smaller than the last) rather than -> 0. Under the sole "
              "recoverable (unrescaled) refinement map, the measure sequence does NOT converge to a coherent limit.")
ws.cell(row=r, column=1).font = Font(bold=True, color="C00000")
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
ws.row_dimensions[r].height = 55

# 98 Weighted Convergence
ws = new_sheet("98 C-004F Convergence")
write_title(ws, "WEIGHTED-SPACE CONVERGENCE TEST (Outcome A-E)", "")
ws["A5"] = f"Outcome: {REG['C004F_7_weighted_convergence_outcome']['outcome']}"
ws["A5"].font = Font(bold=True, size=13, color="C00000")
ws["A7"] = REG["C004F_7_weighted_convergence_outcome"]["justification"]
ws["A7"].font = BODY_FONT; ws["A7"].alignment = WRAP
ws.merge_cells("A7:E7"); ws.row_dimensions[7].height = 60
ws["A9"] = REG["C004F_7_weighted_convergence_outcome"]["caveat_section_17"]
ws["A9"].font = NOTE_FONT; ws["A9"].alignment = WRAP
ws.merge_cells("A9:E9"); ws.row_dimensions[9].height = 40

# 99 Continuum Recovery (blocked)
ws = new_sheet("99 C-004F Continuum Recovery")
write_title(ws, "C-004F.8-10: CONTINUUM DIRICHLET FORM / GENERATOR / FP CORRESPONDENCE", "")
ws["A5"] = REG["C004F_8_9_10"]["status"]
ws["A5"].font = BODY_FONT; ws["A5"].alignment = WRAP
ws.merge_cells("A5:E5"); ws.row_dimensions[5].height = 60

# 100 Fokker-Planck Bridge
ws = new_sheet("100 C-004F FP Bridge")
write_title(ws, "FOKKER-PLANCK BRIDGE STATUS", "")
write_table(ws, 4, ["Object", "Status"], [
    ("Discrete generator Q_k* = -L_kD_k^{-1}", "DERIVED / VERIFIED (C-004D, re-verified this run)"),
    ("Continuum limit of (G_k, mu_k, E_k)", "DISPROVEN for unrescaled map / OPEN for rescaled (untestable)"),
    ("Correspondence with L*P=-div(P grad log P_ss)", "NOT TESTABLE -- blocked upstream"),
], widths=[46, 90])

# 101 Entropy Production
ws = new_sheet("101 C-004F Entropy Production")
write_title(ws, "ENTROPY-PRODUCTION CONSEQUENCE (Section 13)",
            "Does F[P||P_ss] yield sigma=-dF/dt>=0 directly from Q_k*, without choosing P(0) or t?")
ws["A5"] = "YES, as an inequality (not a number)."
ws["A5"].font = Font(bold=True, size=12)
ws["A7"] = REG["entropy_production_consequence_section_13"]["answer"]
ws["A7"].font = BODY_FONT; ws["A7"].alignment = WRAP
ws.merge_cells("A7:E7"); ws.row_dimensions[7].height = 140
r = 10
rows = [(k, v["all_sigma_nonneg"]) for k, v in ENT.items()]
write_table(ws, r, ["Shell", "sigma>=0 across 12 sampled (P(0),t)?"], rows, widths=[8, 34])
r += len(rows) + 2
ws.cell(row=r, column=1,
        value="Status: SIGN DERIVED (confirms/formalizes BRIDGE B-006's own proof method on the concrete Q*). "
              "MAGNITUDE remains PROVEN NON-IDENTIFIABLE (C-004C, not reopened).")
ws.cell(row=r, column=1).font = Font(bold=True)
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
ws.row_dimensions[r].height = 40

# 102 Target Independence
ws = new_sheet("102 C-004F Target Independence")
write_title(ws, "TARGET-INDEPENDENCE AUDIT", "")
ws["A5"] = REG["target_independence_audit"]["result"]
ws["A5"].font = BODY_FONT; ws["A5"].alignment = WRAP
ws.merge_cells("A5:E5"); ws.row_dimensions[5].height = 40
ws["A7"] = REG["target_independence_audit"]["hidden_fitted_constants_check"]
ws["A7"].font = BODY_FONT; ws["A7"].alignment = WRAP
ws.merge_cells("A7:E7"); ws.row_dimensions[7].height = 30

# 103 Theorem Registry
ws = new_sheet("103 C-004F Theorem Registry")
write_title(ws, "C-004F THEOREM REGISTRY", "")
write_table(ws, 4, ["Theorem", "Status"], list(REG["theorem_registry"].items()), widths=[42, 100])

# 104 Closure Matrix
ws = new_sheet("104 C-004F Closure Matrix")
write_title(ws, "C-004F CLOSURE MATRIX", "")
write_table(ws, 4, ["Object", "Status"], [
    ("C-004F.1 Degree measure", REG["recovery_order_status"]["C-004F.1_degree_measure"]),
    ("C-004F.2 Weighted Hilbert space", REG["recovery_order_status"]["C-004F.2_weighted_hilbert_space"]),
    ("C-004F.3 Weighted generator", REG["recovery_order_status"]["C-004F.3_weighted_generator"]),
    ("C-004F.4 Dirichlet form", REG["recovery_order_status"]["C-004F.4_dirichlet_form"]),
    ("C-004F.5 Refinement map", REG["recovery_order_status"]["C-004F.5_refinement_map"]),
    ("C-004F.6 Measure pushforward", REG["recovery_order_status"]["C-004F.6_measure_pushforward"]),
    ("C-004F.7 Weighted convergence", REG["recovery_order_status"]["C-004F.7_weighted_convergence"]),
    ("C-004F.8-10 Continuum/FP", "NOT ATTEMPTED (blocked)"),
    ("THM-C004F.1-3", "PROVEN"),
    ("THM-C004F.4", "DISPROVEN"),
    ("THM-C004F.5", "DISPROVEN (unrescaled) / OPEN (rescaled, untestable)"),
    ("THM-C004F.6", "OPEN / NOT TESTABLE"),
    ("sigma_k sign", "CONFIRMED (SIGN DERIVED) via concrete Q*"),
    ("sigma_k magnitude", "PROVEN NON-IDENTIFIABLE (C-004C, unchanged)"),
    ("G*", "UNRESOLVED (unchanged)"),
], widths=[36, 90])

# 105 Next Dependency
ws = new_sheet("105 C-004F Next Dependency")
write_title(ws, "EXACTLY ONE NEXT UNRESOLVED UPSTREAM DEPENDENCY", "")
ws["A4"] = ("Recover or derive, from admissible source material only (not invented here), the explicit "
            "metric/edge-length RESCALING LAW for the ARBS K_{2^k,2^k}/R_{k-1}xL_k construction specifically "
            "-- the 'global isotropic contraction constant alpha, 0<alpha<1' the whitepaper asserts abstractly "
            "for the generic refinement functor R, but never ties to a concrete formula for this particular "
            "shell construction. Without it, the combinatorial refinement map (fully derived this run: G_k is "
            "an induced-subgraph inclusion into G_k+1, with only R_k vertices gaining exactly 2^(k+1) new edges "
            "each) cannot be turned into a resolution-refining map compatible with any continuum-limit test, and "
            "no rescaled version of the C-004F.6/7 convergence test can be attempted. Under the only recoverable "
            "(unrescaled) map, this run establishes -- with strong geometric numerical evidence over 9 shell "
            "transitions -- that the degree-measure sequence does NOT converge to a coherent limit; that result "
            "stands as the best-supported current answer to C-004F unless and until this rescaling law is "
            "recovered from source material.")
ws["A4"].font = Font(name="Arial", size=12, bold=True)
ws["A4"].alignment = WRAP
ws.merge_cells("A4:E4")
ws.row_dimensions[4].height = 200

wb.save(OUT)
print("saved", OUT)
print("total sheets:", len(wb.sheetnames))
