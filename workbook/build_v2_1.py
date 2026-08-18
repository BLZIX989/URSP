"""
Build v2.1 from v2.0: preserves all 163 existing sheets, appends the
ARBS-GRAPH-REALIZATION-001 sheets.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import json

SRC = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v2.0.xlsx"
OUT = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v2.1.xlsx"
RECON = "/home/user/URSP/reconstruction"

def load(name):
    with open(f"{RECON}/{name}") as f:
        return json.load(f)

CG = load("arbs_canonical_graph.json")
CMP = load("arbs_graph_comparison.json")
ORIG = load("arbs_checkpoint_origin.json")
OBS = load("ARBS-GRAPH-REALIZATION-OBSTRUCTION-001.json")
TI = load("arbs_graph_target_independence.json")

wb = openpyxl.load_workbook(SRC, data_only=False)

TITLE_FONT = Font(name="Arial", size=12, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def new_sheet(name):
    return wb.create_sheet(name)


def write_title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = WRAP


def write_table(ws, start_row, headers, rows, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + 1 + len(rows)


def block(ws, r, title, text, height=60):
    ws.cell(row=r, column=1, value=title).font = Font(bold=True)
    r += 1
    ws.cell(row=r, column=1, value=text)
    ws.cell(row=r, column=1).font = BODY_FONT
    ws.cell(row=r, column=1).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = height
    return r + 2


# 165 Overview
ws = new_sheet("165 ARBS-GRAPH-001 Overview")
write_title(ws, "ARBS-GRAPH-REALIZATION-001", "G_ARBS = canonical shell-1 graph. tilde_G_k = previously reconstructed numerical family. NEVER conflated below.")
r = block(ws, 5, "Bottom line", CG["conclusion"], height=90)
r = block(ws, r, "Stop condition triggered", OBS["trigger"], height=50)

# 166 Source Recovery
ws = new_sheet("166 ARBS-GRAPH-002 Source")
write_title(ws, "EVERY OCCURRENCE OF G=(V,E,W) FOUND IN SOURCE", "")
rows = [(o["location"], o["text"], o["nature"]) for o in CG["every_occurrence_found"]]
write_table(ws, 5, ["Location", "Text (verbatim/paraphrase)", "Nature"], rows, widths=[30, 60, 50])

# 167 Vertex Set
ws = new_sheet("167 ARBS-GRAPH-003 Vertex Set")
write_title(ws, "THM-ARBS-GRAPH-VERTEX-001", "")
r = block(ws, 5, "Recovered", "V = entities/states (semantic definition only). No concrete vertex list, count, or generation rule given anywhere.", height=40)
r = block(ws, r, "Status", "OPEN -- semantic role recovered; concrete instantiation NOT source-supported.", height=30)

# 168 Edge Set
ws = new_sheet("168 ARBS-GRAPH-004 Edge Set")
write_title(ws, "THM-ARBS-GRAPH-EDGE-001", "")
r = block(ws, 5, "Recovered", "E = relations (semantic definition only). A_ij=w_ij (i!=j) given as a GENERIC formula relating weights to an adjacency-like matrix, not a construction of E itself. No concrete edge list or generation rule given.", height=50)
r = block(ws, r, "Status", "OPEN -- semantic role recovered; concrete instantiation NOT source-supported.", height=30)

# 169 Weights
ws = new_sheet("169 ARBS-GRAPH-005 Weights")
write_title(ws, "THM-ARBS-GRAPH-WEIGHT-001", "")
r = block(ws, 5, "Recovered", "W = relation strengths (semantic definition only). No concrete weight formula, normalization, or value range given anywhere in source.", height=40)
r = block(ws, r, "Status", "OPEN -- semantic role recovered; concrete instantiation NOT source-supported.", height=30)

# 170 Bipartition
ws = new_sheet("170 ARBS-GRAPH-006 Bipartition")
write_title(ws, "THM-ARBS-GRAPH-BIPARTITE-001", "")
r = block(ws, 5, "TH-ARBS-001A (Bipartite Reciprocity Lock)", CG["the_two_admissibility_axioms_(concrete, checkable)"]["TH-ARBS-001A_bipartite_reciprocity_lock"], height=40)
r = block(ws, r, "Test against tilde_G_k", CMP["TH-ARBS-001A_bipartite_reciprocity_lock_test"]["test_result"], height=60)
r = block(ws, r, "Status", CMP["TH-ARBS-001A_bipartite_reciprocity_lock_test"]["status"], height=30)

# 171 Dependency
ws = new_sheet("171 ARBS-GRAPH-007 Dependency")
write_title(ws, "THM-ARBS-GRAPH-DEPENDENCY-001", "")
r = block(ws, 5, "TH-ARBS-001B (Shell Nilpotency Lock)", CG["the_two_admissibility_axioms_(concrete, checkable)"]["TH-ARBS-001B_shell_nilpotency_lock"], height=50)
r = block(ws, r, "Test performed", CMP["TH-ARBS-001B_shell_nilpotency_lock_test"]["test_performed"], height=60)
r = block(ws, r, "Numeric result", CMP["TH-ARBS-001B_shell_nilpotency_lock_test"]["numeric_result"], height=50)
r = block(ws, r, "Status", CMP["TH-ARBS-001B_shell_nilpotency_lock_test"]["status"], height=30)

# 172 Hypergraph Relation
ws = new_sheet("172 ARBS-GRAPH-008 Hypergraph")
write_title(ws, "THM-ARBS-GRAPH-HYPERGRAPH-001", "")
r = block(ws, 5, "Finding", "No map P_G: H_ARBS -> G_ARBS is defined anywhere (consistent with ARBS-ONTOLOGY-RECOVERY-001's finding that no ARBS->tilde_G_k map exists at all). The hypergraph-to-graph relation remains entirely unspecified.", height=70)
r = block(ws, r, "Status", "OPEN", height=30)

# 173 Canonical Graph (obstruction)
ws = new_sheet("173 ARBS-GRAPH-009 Canonical")
write_title(ws, "CANONICAL GRAPH RECOVERY -- OBSTRUCTION", "")
r = block(ws, 5, "Conclusion", CG["conclusion"], height=90)
r = block(ws, r, "Category found (not a unique graph)", CG["every_occurrence_found"][3]["text"], height=50)
r = block(ws, r, "Artifacts NOT produced", CG["no_canonical_graph_artifacts_produced"], height=50)

# 174 Adjacency (N/A)
ws = new_sheet("174 ARBS-GRAPH-010 Adjacency")
write_title(ws, "CANONICAL ADJACENCY -- NOT APPLICABLE", "No G_ARBS exists to construct an adjacency matrix from. See sheet 173.")

# 175 Laplacian
ws = new_sheet("175 ARBS-GRAPH-011 Laplacian")
write_title(ws, "THM-ARBS-GRAPH-LAPLACIAN-001", "")
r = block(ws, 5, "Finding", "BRIDGE B-001 states L=D-A is the Laplacian 'given G=(V,E,W)' -- a general fact about ANY graph with this schema, not a claim that ARBS defines a special hypergraph-native operator. No L_ARBS distinct from the standard L=D-A construction was found. This is consistent with (not evidence against) tilde_L_k=D_k-A_k remaining the correct operator FOR tilde_G_k specifically -- but does not establish tilde_L_k = L_ARBS since G_ARBS itself is unresolved.", height=110)
r = block(ws, r, "Status", "L=D-A is the certified GENERAL formula (BRIDGE B-001); no ARBS-specific alternative found; applicability to a specific G_ARBS remains OPEN pending sheet 173's obstruction.", height=50)

# 176 Comparison with tilde-G
ws = new_sheet("176 ARBS-GRAPH-012 Comparison")
write_title(ws, "STRUCTURAL COMPARISON: ADMISSIBILITY, NOT IDENTITY", CMP["note"])
r = block(ws, 5, "Conclusion on admissibility", CMP["conclusion_on_admissibility"], height=120)

# 177 Structural Invariants
ws = new_sheet("177 ARBS-GRAPH-013 Invariants")
write_title(ws, "STRUCTURAL INVARIANTS TESTED", "")
write_table(ws, 5, ["Invariant", "tilde_G_k result"], [
    ("Global bipartite 2-coloring (TH-ARBS-001A)", "SATISFIED, all tested shells"),
    ("Directed nilpotency under construction-order orientation (TH-ARBS-001B)", "SATISFIED, nilpotency index = 2k+1 for G_k"),
    ("Hyperedges (arity>=3)", "NONE present (G_k is a simple graph)"),
    ("Node/edge typing", "NONE present (G_k is untyped)"),
], widths=[50, 60])

# 178 Spectral Comparison
ws = new_sheet("178 ARBS-GRAPH-014 Spectral")
write_title(ws, "SPECTRAL COMPARISON -- NOT ATTEMPTED", "Per the governing proof order (structure before spectrum), and since no G_ARBS/L_ARBS exists to compute a spectrum from, no Spec(L_ARBS) vs Spec(tilde_L_k) comparison is performed this run. All spectral facts about tilde_G_k established in runs 1-9 stand, unchanged, as tilde_G_k-level results.")

# 179 Checkpoint Origin
ws = new_sheet("179 ARBS-GRAPH-015 Checkpoint")
write_title(ws, "ARBS-CHECKPOINT-ORIGIN-001", "")
write_table(ws, 5, ["Option", "Description"], list(ORIG["options"].items()), widths=[10, 100])
r = 11
r = write_table(ws, r, ["Evidence"], [[x] for x in ORIG["evidence"]], widths=[140])
r = block(ws, r, "Assessment", ORIG["assessment"], height=140)
r = block(ws, r, "Status", ORIG["status"], height=50)

# 180 Prior Result Reclassification
ws = new_sheet("180 ARBS-GRAPH-016 Reclassify")
write_title(ws, "PRIOR RESULT RECLASSIFICATION (UPDATE)", "Unchanged from ARBS-ONTOLOGY-RECOVERY-001's table (workbook sheet 158) -- this run adds no new invalidation, only sharper reasoning for WHY the ARBS-level interpretation remains open (concrete V,E,W obstruction) plus positive admissibility evidence not previously available.")
r = block(ws, 5, "Net effect on prior reclassification table", "All entries in sheet 158 stand as written. New nuance added: tilde_G_k is now known to be an ADMISSIBLE member of the Graph_ARBS category (satisfies TH-ARBS-001A/B), which is stronger evidence of compatibility than run 10 had -- but still short of identity/uniqueness, since G_ARBS itself remains unspecified.", height=70)

# 181 Target Independence
ws = new_sheet("181 ARBS-GRAPH-017 Target Indep")
write_title(ws, "TARGET-INDEPENDENCE AUDIT", "")
r = block(ws, 5, "Checked", ", ".join(TI["checked"]), height=40)
r = block(ws, r, "Result", TI["result"], height=60)

# 182 Theorem Registry
ws = new_sheet("182 ARBS-GRAPH-018 Theorems")
write_title(ws, "THEOREM REGISTRY", "")
write_table(ws, 5, ["Theorem", "Status"], [
    ("THM-ARBS-GRAPH-VERTEX-001", "OPEN -- semantic role recovered, concrete V not source-supported"),
    ("THM-ARBS-GRAPH-EDGE-001", "OPEN -- semantic role recovered, concrete E not source-supported"),
    ("THM-ARBS-GRAPH-WEIGHT-001", "OPEN -- semantic role recovered, concrete W not source-supported"),
    ("THM-ARBS-GRAPH-BIPARTITE-001", "tilde_G_k SATISFIES TH-ARBS-001A (VERIFIED); uniqueness OPEN"),
    ("THM-ARBS-GRAPH-DEPENDENCY-001", "tilde_G_k SATISFIES TH-ARBS-001B under natural orientation (VERIFIED); uniqueness OPEN"),
    ("THM-ARBS-GRAPH-HYPERGRAPH-001", "OPEN -- no hypergraph-to-graph map found"),
    ("THM-ARBS-GRAPH-LAPLACIAN-001", "L=D-A CERTIFIED as the general formula (BRIDGE B-001); ARBS-specific alternative NOT FOUND; applicability OPEN"),
    ("THM-ARBS-GRAPH-REALIZATION-001", "OBSTRUCTION -- G_ARBS cannot be uniquely reconstructed from source"),
    ("THM-ARBS-CHECKPOINT-ORIGIN-001", "CALCULATED/ASSESSED (blend of D and E); not provable to certainty"),
    ("THM-ARBS-SPECTRAL-INHERITANCE-001", "NOT ATTEMPTED -- no G_ARBS/L_ARBS to compare against"),
], widths=[42, 100])

# 183 Falsification Ledger
ws = new_sheet("183 ARBS-GRAPH-019 Falsif")
write_title(ws, "FALSIFICATION LEDGER", "")
write_table(ws, 5, ["Claim tested", "Result"], [
    ("G=(V,E,W) is concretely defined anywhere in source", "FALSIFIED -- only semantic/schema-level definitions found"),
    ("Graph_ARBS names a single canonical graph", "FALSIFIED -- it names a CATEGORY of admissible graphs (2 axioms)"),
    ("The SIT Delta->G construction reproduces a bipartite-shell graph", "FALSIFIED, per the source's own explicit test ('converges to near-complete graph, not the claimed Recursive Bipartite Shell Graph')"),
    ("tilde_G_k is an arbitrary, structurally unrelated graph", "FALSIFIED -- it satisfies both concrete ARBS admissibility axioms"),
    ("tilde_G_k is proven identical/realization-equal to G_ARBS", "NOT ESTABLISHED (no G_ARBS exists to test against) -- distinct from 'falsified'"),
], widths=[60, 90])

# 184 MDCL
ws = new_sheet("184 ARBS-GRAPH-020 MDCL")
write_title(ws, "CORRECTED MDCL", "")
ws["A5"] = ("Gamma < ARBS < ARBS shell 1 < G_ARBS [OBSTRUCTION -- only a CATEGORY (Graph_ARBS, 2 admissibility axioms) "
            "is specified, not a unique graph] < L_ARBS [OPEN] < Spec(L_ARBS) [OPEN]\n\n"
            "SEPARATELY (unchanged, fully valid):\n"
            "tilde_G_k [ADMISSIBLE member of Graph_ARBS category, VERIFIED] < tilde_L_k=D_k-A_k < Spec(tilde_L_k) "
            "< [everything computed in runs 1-9]\n\n"
            "No proven realization map P: G_ARBS -> tilde_G_k connects the two chains, because G_ARBS is not "
            "concretely specified for such a map's domain to even be defined.")
ws["A5"].font = BODY_FONT
ws["A5"].alignment = WRAP
ws.merge_cells("A5:F5")
ws.row_dimensions[5].height = 160

# 185 Closure Matrix
ws = new_sheet("185 ARBS-GRAPH-021 Closure")
write_title(ws, "CLOSURE MATRIX", "")
write_table(ws, 5, ["Object", "Status"], [
    ("G_ARBS = (V,E,W) concrete instantiation", "OBSTRUCTION -- not source-supported"),
    ("Graph_ARBS category + 2 admissibility axioms", "DERIVED / VERIFIED (documented, checkable)"),
    ("tilde_G_k satisfies TH-ARBS-001A", "VERIFIED"),
    ("tilde_G_k satisfies TH-ARBS-001B", "VERIFIED"),
    ("tilde_G_k = G_ARBS (identity)", "NOT ESTABLISHED -- no target to compare against"),
    ("SIT Delta->G alternative construction", "REVIEWED -- unfixed parameters, explicitly fails to reproduce bipartite-shell structure per source's own test"),
    ("Checkpoint origin", "ASSESSED (blend of Outcomes D/E); not provable to certainty from available source"),
    ("All runs 1-9 numerical results", "PRESERVED, VALID as tilde_G_k-level facts"),
    ("sigma_k, Pi_0, G*, NCG, flavor, scale", "UNTOUCHED, per stop condition"),
    ("PACKET-REFINEMENT-RECOVERY-001", "NOT ACTIVATED"),
], widths=[46, 100])

# 186 Next Dependency
ws = new_sheet("186 ARBS-GRAPH-022 Next Dep")
write_title(ws, "EXACTLY ONE NEXT UNRESOLVED UPSTREAM DEPENDENCY", "")
ws["A4"] = OBS["next_dependency"]
ws["A4"].font = Font(name="Arial", size=12, bold=True)
ws["A4"].alignment = WRAP
ws.merge_cells("A4:F4")
ws.row_dimensions[4].height = 200

wb.save(OUT)
print("saved", OUT)
print("total sheets:", len(wb.sheetnames))
