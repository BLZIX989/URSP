"""
Build v1.7 from v1.6: preserves all 105 existing sheets, appends the C-004G
edge-length rescaling law recovery sheets, with the recovered equation
images embedded directly.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
import json

SRC = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v1.6.xlsx"
OUT = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v1.7.xlsx"
RECON = "/home/user/URSP/reconstruction"

with open(f"{RECON}/c004g_edge_scaling.json") as f:
    REG = json.load(f)
with open(f"{RECON}/c004g_diameter.json") as f:
    DIAM = json.load(f)

wb = openpyxl.load_workbook(SRC, data_only=False)

TITLE_FONT = Font(name="Arial", size=12, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def new_sheet(name):
    return wb.create_sheet(name)


def write_title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = WRAP


def write_table(ws, start_row, headers, rows, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + 1 + len(rows)


# ---------------------------------------------------------------------
# 106 C-004G Recovered Law
# ---------------------------------------------------------------------
ws = new_sheet("106 C-004G Recovered Law")
write_title(ws, "C-004G -- ARBS EDGE-LENGTH RESCALING LAW, RECOVERED",
            "Recovered from source/Combined_Compiler_Theories_Whitepaper.docx, Part I Section 2.1. This law "
            "was silently absent from the prior text-extraction corpus (54_SOURCE_TEXT_CORPUS) because it is "
            "embedded as RENDERED EQUATION IMAGES in the original .docx, not text or OMML markup -- invisible "
            "to a paragraph-text scraper. Recovered by unzipping the .docx directly and reading the three "
            "embedded PNG images anchored to the relevant bullets.")

r = 5
ws.cell(row=r, column=1, value="Primal Edge Scaling").font = Font(bold=True, size=11)
r += 1
img1 = XLImage(f"{RECON}/equations_recovered/primal_edge_scaling.png")
img1.anchor = f"A{r}"
ws.add_image(img1)
r += 6
ws.cell(row=r, column=1, value="Meaning: every child edge e' produced by refinement functor R from parent edge e "
                                "has length = alpha * (parent length); alpha is a fixed global isotropic contraction "
                                "constant, 0<alpha<1.")
ws.cell(row=r, column=1).font = BODY_FONT
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
ws.row_dimensions[r].height = 40
r += 3

ws.cell(row=r, column=1, value="Induced Metric Invariance").font = Font(bold=True, size=11)
r += 1
img2 = XLImage(f"{RECON}/equations_recovered/induced_metric_invariance.png")
img2.anchor = f"A{r}"
ws.add_image(img2)
r += 6
ws.cell(row=r, column=1, value="Meaning: induced path distance between vertex IMAGES under R is exactly preserved -- "
                                "edges shrink by alpha while path edge-counts grow proportionally, invariant total length.")
ws.cell(row=r, column=1).font = BODY_FONT
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
ws.row_dimensions[r].height = 40
r += 3

ws.cell(row=r, column=1, value="Diameter Stability").font = Font(bold=True, size=11)
r += 1
img3 = XLImage(f"{RECON}/equations_recovered/diameter_stability.png")
img3.anchor = f"A{r}"
ws.add_image(img3)
r += 6
ws.cell(row=r, column=1, value="Meaning: the global diameter of the metric space stays bounded by a fixed D_max "
                                "as n increases -- the sequence is an increasingly FINE approximation of a FIXED, "
                                "bounded space.")
ws.cell(row=r, column=1).font = BODY_FONT
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
ws.row_dimensions[r].height = 40
r += 3
ws.cell(row=r, column=1, value=f"Numeric value of alpha: {REG['recovered_law']['numeric_value_of_alpha']}")
ws.cell(row=r, column=1).font = NOTE_FONT
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
ws.row_dimensions[r].height = 40
r += 2
ws.cell(row=r, column=1, value=f"Connection to ARBS specifically: {REG['recovered_law']['connection_to_ARBS_specifically']}")
ws.cell(row=r, column=1).font = NOTE_FONT
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
ws.row_dimensions[r].height = 55

# ---------------------------------------------------------------------
# 107 C-004G Compatibility Tests
# ---------------------------------------------------------------------
ws = new_sheet("107 C-004G Compatibility")
write_title(ws, "COMPATIBILITY TEST: DOES THIS LAW APPLY TO THE ARBS SHELL CONSTRUCTION?", "")
tests = REG["compatibility_test_against_ARBS"]
write_table(ws, 4, ["Test", "Finding", "Conclusion"], [
    ("1. Edge-subdivision domain", tests["test_1_no_edge_subdivision_domain"]["finding"], tests["test_1_no_edge_subdivision_domain"]["conclusion"]),
    ("2. Diameter Stability", tests["test_2_diameter_stability_directly_falsified"]["finding"], tests["test_2_diameter_stability_directly_falsified"]["conclusion"]),
    ("3. Forced alpha value", tests["test_3_forced_alpha_would_be_1_not_less_than_1"]["finding"], tests["test_3_forced_alpha_would_be_1_not_less_than_1"]["conclusion"]),
], widths=[26, 80, 40])

# ---------------------------------------------------------------------
# 108 C-004G Diameter Computation
# ---------------------------------------------------------------------
ws = new_sheet("108 C-004G Diameter Calc")
write_title(ws, "EXACT GRAPH DIAMETER, G0-G7", "Unweighted shortest-path hop count. Direct test of Diam(M_n,d_n)<=D_max<infinity.")
rows = [(k, v["N"], v["diameter_hops"], v["predicted_2n+1"], v["matches_prediction"]) for k, v in DIAM.items()]
write_table(ws, 4, ["Shell", "N", "Diameter (hops)", "Predicted 2n+1", "Matches?"], rows, widths=[8, 8, 16, 16, 12])
r = 4 + len(rows) + 2
ws.cell(row=r, column=1, value="Diameter grows exactly as 2n+1 -- linear, unbounded. Diameter Stability is FALSIFIED "
                                "by direct counterexample, not merely unverified.")
ws.cell(row=r, column=1).font = Font(bold=True, color="C00000")
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)

# ---------------------------------------------------------------------
# 109 C-004G Closure
# ---------------------------------------------------------------------
ws = new_sheet("109 C-004G Closure")
write_title(ws, "C-004G FINAL CLOSURE", "")
ws["A5"] = REG["conclusion"]["statement"]
ws["A5"].font = BODY_FONT; ws["A5"].alignment = WRAP
ws.merge_cells("A5:E5"); ws.row_dimensions[5].height = 110
r = 8
ws.cell(row=r, column=1, value="Consequence for C-004F").font = Font(bold=True)
r += 1
ws.cell(row=r, column=1, value=REG["conclusion"]["consequence_for_C004F"])
ws.cell(row=r, column=1).font = BODY_FONT; ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5); ws.row_dimensions[r].height = 70
r += 3
ws.cell(row=r, column=1, value="Consequence for RF-001..005 / BRIDGE B-004").font = Font(bold=True)
r += 1
ws.cell(row=r, column=1, value=REG["conclusion"]["consequence_for_RF001_005"])
ws.cell(row=r, column=1).font = BODY_FONT; ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5); ws.row_dimensions[r].height = 60
r += 3
write_table(ws, r, ["Object", "Status"], list(REG["status"].items()), widths=[36, 90])

wb.save(OUT)
print("saved", OUT)
print("total sheets:", len(wb.sheetnames))
