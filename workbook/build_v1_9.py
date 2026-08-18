"""
Build v1.9 from v1.8: preserves all 121 existing sheets, appends the C-004I
diffusion-metric convergence and measure recovery sheets.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import json

SRC = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v1.8.xlsx"
OUT = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v1.9.xlsx"
RECON = "/home/user/URSP/reconstruction"

def load(name):
    with open(f"{RECON}/{name}") as f:
        return json.load(f)

MET = load("C004I_metric_registry.json")
REF = load("C004I_refinement_registry.json")
MEA = load("C004I_measure_registry.json")
CONV = load("C004I_convergence_registry.json")
SYN = load("C004I_synthesis.json")

wb = openpyxl.load_workbook(SRC, data_only=False)

TITLE_FONT = Font(name="Arial", size=12, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def new_sheet(name):
    return wb.create_sheet(name)


def write_title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = WRAP


def write_table(ws, start_row, headers, rows, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + 1 + len(rows)


def block(ws, r, title, text, height=60):
    ws.cell(row=r, column=1, value=title).font = Font(bold=True)
    r += 1
    ws.cell(row=r, column=1, value=text)
    ws.cell(row=r, column=1).font = BODY_FONT
    ws.cell(row=r, column=1).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = height
    return r + 2


# 122 Overview
ws = new_sheet("122 C-004I Overview")
write_title(ws, "C-004I -- ARBS DIFFUSION-METRIC CONVERGENCE AND MEASURE RECOVERY",
            "C-004H found the diffusion-distance diameter bounded; this run tests whether ARBS admits a "
            "genuine metric-measure continuum limit under this source-certified metric.")
r = block(ws, 5, "Headline finding", "The Wasserstein-1 distance (using the diffusion metric as ground cost) "
          "between the pushforward degree measure and the next shell's measure decays geometrically toward 0, "
          "converging to the SAME sqrt(2) rate found for metric-invariance violation -- even though total "
          "variation distance (the norm implicitly used in C-004F/H) does not. This is a genuinely new, "
          "positive, cross-validated result -- not a full proof.", height=70)

# 123 Source Metric
ws = new_sheet("123 C-004I Source Metric")
write_title(ws, "C-004I.1 -- EXACT SOURCE DEFINITION OF DIFFUSION DISTANCE", "")
d = MET["source_definition"]
r = block(ws, 5, "Formula", d["formula"])
r = block(ws, r, "Source", d["source"], height=50)
r = block(ws, r, "No time parameter", d["no_time_parameter"], height=40)
r = block(ws, r, "Normalization", d["normalization"], height=30)

# 124 Diffusion Distance (proof)
ws = new_sheet("124 C-004I Diffusion Distance")
write_title(ws, "d(i,j)^2 = R_eff(i,j) -- ANALYTIC PROOF", "")
p = MET["identity_proof_d2_equals_Reff"]
r = write_table(ws, 5, ["Step"], [[s] for s in p["proof"]], widths=[140])
r = block(ws, r, "Status", p["status"])
r = block(ws, r, "Numerical confirmation", p["numerical_confirmation"], height=40)

# 125 Effective Resistance
ws = new_sheet("125 C-004I Eff Resistance")
write_title(ws, "EFFECTIVE RESISTANCE -- IDENTITY CHECK", "max|d(i,j)^2 - R_eff(i,j)| across G0-G8")
rows = [(k, v) for k, v in CONV["identity_check_d2_eq_Reff"].items()]
write_table(ws, 5, ["Shell", "max|d^2-R_eff|"], rows, widths=[10, 20])

# 126 Closed Form
ws = new_sheet("126 C-004I Closed Form")
write_title(ws, "C-004I.2 -- CLOSED-FORM ANALYSIS", "")
cf = MET["closed_form_analysis"]
r = block(ws, 5, "Vertex-class symmetry (claim + proof sketch)",
          cf["vertex_class_symmetry"]["claim"] + " " + cf["vertex_class_symmetry"]["proof_sketch"], height=90)
r = block(ws, r, "Status", cf["vertex_class_symmetry"]["status"], height=40)
rows = [(k, v) for k, v in CONV["vertex_class_symmetry_max_std"].items()]
r = write_table(ws, r, ["Shell", "max std-dev within class-pair (should be ~0)"], rows, widths=[10, 40])
r += 2
r = block(ws, r, "Quotient matrix reduction", cf["quotient_matrix_reduction"]["description"], height=70)
r = block(ws, r, "Status", cf["quotient_matrix_reduction"]["status"], height=50)
r = block(ws, r, "D_inf", cf["quotient_matrix_reduction"]["diam_D_inf"], height=40)

# 127 Refinement Map
ws = new_sheet("127 C-004I Refinement Map")
write_title(ws, "C-004I.3 -- REFINEMENT MAP", "")
rm = REF["refinement_map"]
r = block(ws, 5, "rho_k", rm["rho_k"], height=40)
r = block(ws, r, "Distinction preserved", rm["distinction_preserved"], height=50)

# 128 Metric Invariance (epsilon_k precise)
ws = new_sheet("128 C-004I Metric Invariance")
write_title(ws, "C-004I.3/4 -- epsilon_k PRECISE DEFINITION AND METRIC INVARIANCE", "")
ek = REF["epsilon_k_precise_definition"]
r = block(ws, 5, "Precise definition", ek["definition"], height=50)
r = block(ws, r, "Why this quantity", ek["why_this_quantity"], height=40)
rows = [(k, v) for k, v in CONV["epsilon_abs_metric_invariance_violation"].items()]
r = write_table(ws, r, ["Transition", "epsilon_k = max|d_k+1(i,j)-d_k(i,j)|"], rows, widths=[16, 30])
r += 1
ws.cell(row=r, column=1, value="Ratios epsilon_k/epsilon_k+1 (all consecutive transitions from G2->G3):").font = Font(bold=True)
r += 1
ws.cell(row=r, column=1, value=str(CONV["epsilon_abs_ratios"]))
ws.cell(row=r, column=1).font = BODY_FONT
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
ws.row_dimensions[r].height = 40
r += 2
r = block(ws, r, "Classification", REF["metric_invariance_classification"]["answer"], height=60)
r = block(ws, r, "Rescaling needed?", REF["metric_invariance_classification"]["no_rescaling_needed"], height=50)

# 129 Error Scaling
ws = new_sheet("129 C-004I Error Scaling")
write_title(ws, "SQUARED-DISTANCE VIOLATION (alternative epsilon definition, cross-check)", "")
rows = [(k, v) for k, v in CONV["epsilon_sq_metric_invariance_violation"].items()]
write_table(ws, 5, ["Transition", "max|d_k+1^2 - d_k^2|"], rows, widths=[16, 30])

# 130 Degree Measure
ws = new_sheet("130 C-004I Degree Measure")
write_title(ws, "C-004I.5 -- DEGREE MEASURE DEFECT, PRECISE DEFINITION", "")
md = MEA["measure_defect_precise_definition"]
for key in ["TV", "L1", "L2", "Wasserstein1"]:
    ws.cell(row=5 + list(["TV","L1","L2","Wasserstein1"]).index(key)*2, column=1, value=f"{key}: {md[key]}").font = BODY_FONT
r = 14
r = block(ws, r, "TV/L1/L2 trend", md["numeric_trend_TV_L1_L2"], height=70)
r = block(ws, r, "TV/L1/L2 verdict", md["TV_L1_L2_verdict"], height=40)

# 131 Measure Pushforward (recomputed values table)
ws = new_sheet("131 C-004I Measure Pushforward")
write_title(ws, "MEASURE DEFECT VALUES, G0-G10", "")
rows = [(k, v["TV"], v["L1"], v["L2"], CONV["wasserstein1_diffusion_ground_metric"].get(k))
        for k, v in CONV["measure_defect_TV_L1_L2"].items()]
write_table(ws, 5, ["Transition", "TV", "L1", "L2", "Wasserstein-1 (diffusion)"], rows, widths=[16, 14, 14, 14, 24])

# 132 Measure Defect (Wasserstein headline)
ws = new_sheet("132 C-004I Wasserstein")
write_title(ws, "HEADLINE FINDING: WASSERSTEIN-1 CONVERGENCE", "")
w = MEA["wasserstein1_result"]
r = block(ws, 5, "Finding", w["critical_new_finding"], height=70)
r = block(ws, r, "Mechanism", w["mechanism"], height=90)
r = block(ws, r, "Relation to metric-invariance finding", w["relation_to_metric_invariance_finding"], height=50)
r = block(ws, r, "Status", w["status"], height=80)

# 133 Measure Renormalization
ws = new_sheet("133 C-004I Renormalization")
write_title(ws, "C-004I.6 -- MEASURE RENORMALIZATION SEARCH", "")
rn = MEA["renormalization_search"]
r = block(ws, 5, "Finding", rn["finding"], height=60)
r = block(ws, r, "Search performed", rn["search_performed_anyway"], height=40)
r = block(ws, r, "Status", rn["status"], height=30)

# 134 Metric-Measure Convergence / Outcome
ws = new_sheet("134 C-004I Convergence")
write_title(ws, "C-004I.7/13 -- METRIC-MEASURE CONVERGENCE AND OUTCOME SELECTION", "")
osel = SYN["outcome_selection_C004I_13"]
write_table(ws, 5, ["Option", "Description"], list(osel["options"].items()), widths=[10, 100])
r = 11
ws.cell(row=r, column=1, value="Detail").font = Font(bold=True)
r += 1
r = write_table(ws, r, ["Component finding"], [[x] for x in osel["detail"]], widths=[140])
r = block(ws, r, "Closest honest label", osel["closest_honest_label"], height=110)
r = block(ws, r, "L_k vs Q_k caveat (CRITICAL)", osel["L_k_vs_Q_k_caveat"], height=50)

# 135 Dirichlet Form
ws = new_sheet("135 C-004I Dirichlet Form")
write_title(ws, "C-004I.9 -- DIRICHLET FORM RECONCILIATION (L_k vs Q_k)", "")
df = SYN["dirichlet_form_reconciliation"]
r = block(ws, 5, "Requested form", df["requested_form"], height=50)
r = block(ws, r, "Critical distinction preserved", df["critical_distinction_preserved"], height=80)
r = block(ws, r, "Consequence", df["consequence"], height=90)
r = block(ws, r, "Status", df["status"], height=30)

# 136 Continuum Dimension
ws = new_sheet("136 C-004I Continuum Dimension")
write_title(ws, "C-004I.8 -- CONTINUUM DIMENSION DIAGNOSTICS", "")
cd = SYN["continuum_dimension"]
r = block(ws, 5, "Volume growth", cd["volume_growth_diagnostic"]["finding"], height=90)
r = block(ws, r, "Volume growth status", cd["volume_growth_diagnostic"]["status"], height=30)
r = block(ws, r, "Spectral dimension", cd["spectral_dimension_diagnostic"]["finding"], height=60)
r = block(ws, r, "Spectral dimension status", cd["spectral_dimension_diagnostic"]["status"], height=40)

# 137 Target Independence
ws = new_sheet("137 C-004I Target Independence")
write_title(ws, "TARGET-INDEPENDENCE AUDIT", "")
ti = SYN["target_independence_audit"]
r = block(ws, 5, "Checked objects", ", ".join(ti["checked"]), height=40)
r = block(ws, r, "Result", ti["result"], height=50)

# 138 Falsification Ledger
ws = new_sheet("138 C-004I Falsification Ledger")
write_title(ws, "FALSIFICATION LEDGER ADDITIONS", "")
r = block(ws, 5, "TV as the relevant convergence norm", SYN["falsification_ledger_additions"]["TV_as_the_relevant_convergence_norm"], height=70)
r = block(ws, r, "Edge-scaling law revisit", REF["edge_scaling_law_revisit"]["status"], height=40)
r = block(ws, r, "Any legitimate reinterpretation of R(e) found for ARBS?", REF["edge_scaling_law_revisit"]["any_legitimate_reinterpretation_found"], height=50)

# 139 Closure Matrix
ws = new_sheet("139 C-004I Closure Matrix")
write_title(ws, "C-004I CLOSURE MATRIX", "")
write_table(ws, 5, ["Object", "Status"], [
    ("d(i,j)^2 = R_eff(i,j)", "PROVEN ANALYTICALLY (general graph identity)"),
    ("Vertex-class symmetry", "PROVEN (automorphism argument) + verified to machine precision"),
    ("Quotient-matrix reduction", "STRUCTURE IDENTIFIED; full symbolic diagonalization OPEN"),
    ("D_inf (diameter limit)", "NUMERICALLY CHARACTERIZED (~1.4288690166); not derived in closed form"),
    ("epsilon_k (metric invariance violation)", "PRECISELY DEFINED; ratio = sqrt(2) to 14+ sig figs, 8 consecutive transitions"),
    ("Metric invariance", "ASYMPTOTIC, no rescaling needed (strong numerical evidence, not proven)"),
    ("Edge-scaling law (Primal Edge Scaling)", "PROVEN INCOMPATIBLE (unchanged)"),
    ("TV/L1/L2 measure convergence", "PROVEN NON-CONVERGENT (unchanged)"),
    ("Wasserstein-1 measure convergence", "STRONG NUMERICAL EVIDENCE OF CONVERGENCE (new, headline finding)"),
    ("Dirichlet form (mu-normalized)", "DERIVED/VERIFIED (unchanged, C-004F.4)"),
    ("L_k vs Q_k", "DISTINCT, preserved -- geometric findings pertain to L_k only"),
    ("Continuum dimension", "INCONCLUSIVE (volume growth); PRELIMINARY/unvalidated (~1.8 spectral)"),
    ("Outcome A/B/C/D", "NONE cleanly applies; graded componentwise result recorded instead"),
    ("sigma_k", "PROVEN NON-IDENTIFIABLE (untouched, per instruction)"),
    ("G*", "UNRESOLVED (unchanged)"),
    ("PACKET-REFINEMENT-RECOVERY-001", "NOT ACTIVATED -- ARBS not shown to fail continuum substrate requirements"),
], widths=[40, 90])

# 140 MDCL
ws = new_sheet("140 C-004I MDCL")
write_title(ws, "UPDATED CANONICAL DEPENDENCY CHAIN", "")
ws["A5"] = ("Gamma < ARBS < G_k < A_k,D_k < L_k < Spec(L_k) < [diffusion metric d_k (L_k-based, CONDITIONAL/PROMISING) "
            "< Wasserstein measure convergence (L_k-based, CONDITIONAL/PROMISING)]  -- geometric branch, evidence-graded, not proven\n\n"
            "L_k < Spec(L_k) < t_H,k < lambda_c,k  -- spectral/persistence branch, unaffected\n\n"
            "L_k < D_k^{-1}L_k < Q_k < Q_k* < P_ss,k < F_k < sigma_k (PROVEN NON-IDENTIFIABLE)  -- thermodynamic branch, "
            "UNTOUCHED and NOT resolved by the geometric branch's findings (THM-GEN-DISTINCT-001 keeps these separate)\n\n"
            "Both branches still required to reunite at Pi_0(G_k) < G* -- neither reunification step is attempted this run.")
ws["A5"].font = BODY_FONT
ws["A5"].alignment = WRAP
ws.merge_cells("A5:F5")
ws.row_dimensions[5].height = 160

# 141 Next Dependency
ws = new_sheet("141 C-004I Next Dependency")
write_title(ws, "EXACTLY ONE NEXT UNRESOLVED UPSTREAM DEPENDENCY", "")
ws["A4"] = SYN["next_dependency"]
ws["A4"].font = Font(name="Arial", size=12, bold=True)
ws["A4"].alignment = WRAP
ws.merge_cells("A4:F4")
ws.row_dimensions[4].height = 220

wb.save(OUT)
print("saved", OUT)
print("total sheets:", len(wb.sheetnames))
