"""
Build v2.6 from v2.5: preserves all 247 existing sheets, appends the
UOC-TOE-MASTER-CLOSURE-001 sheets.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import json

SRC = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v2.5.xlsx"
OUT = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v2.6.xlsx"
RECON = "/home/user/URSP/reconstruction"

def load(name):
    with open(f"{RECON}/{name}") as f:
        return json.load(f)

N5 = load("uoc_n5_bipartite_extension.json")
RECONCILE = load("reconciliation_registry.json")
PSG = load("project_state_graph.json")
SEEDREG = load("universal_seed_registry.json")
COMPREG = load("compiler_registry.json")

wb = openpyxl.load_workbook(SRC, data_only=False)

TITLE_FONT = Font(name="Arial", size=12, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")

def new_sheet(name):
    return wb.create_sheet(name)

def write_title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = WRAP

def write_table(ws, start_row, headers, rows, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + 1 + len(rows)

def block(ws, r, title, text, height=60):
    ws.cell(row=r, column=1, value=title).font = Font(bold=True)
    r += 1
    ws.cell(row=r, column=1, value=text)
    ws.cell(row=r, column=1).font = BODY_FONT
    ws.cell(row=r, column=1).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = height
    return r + 2


# 248 N=5 Extension
ws = new_sheet("248 REBUILD-001 N5 Extension")
write_title(ws, "N=5 BIPARTITE SEED EXTENSION", "Exhaustive over the bipartite subset, not sampled")
write_table(ws, 5, ["N", "|F_N^derived_v2|"], [(2,1),(3,4),(4,23),(5,N5["gamma_fixed_bipartite_isomorphism_classes"])], widths=[10,20])
r = 11
r = block(ws, r, "Growth ratios", "4.0, 5.75, 10.04 -- accelerating, not converging", height=30)
r = block(ws, r, "All canonical orientations nilpotent?", str(N5["all_canonical_orientations_nilpotent"]), height=30)
r = block(ws, r, "Symmetric overlap", f'{N5["n_symmetric"]}/{N5["gamma_fixed_bipartite_isomorphism_classes"]} -- still zero', height=30)

# 249 Gauge Reconciliation
ws = new_sheet("249 REBUILD-002 Gauge Reconcile")
write_title(ws, "GAUGE GROUP RECONCILIATION NODE (RECON-001)", "")
node = RECONCILE["nodes"][0]
r = block(ws, 5, "Conflict", node["conflict"], height=90)
r = block(ws, r, "Route A", node["route_vs_result_analysis"]["Route_A"], height=50)
r = block(ws, r, "Route A falsified?", node["route_vs_result_analysis"]["Route_A_falsified"], height=90)
r = block(ws, r, "What this establishes", node["route_vs_result_analysis"]["what_this_establishes"], height=40)
r = block(ws, r, "Does this falsify G_SM itself?", node["route_vs_result_analysis"]["does_this_establish_NOT(G_SM)"], height=50)
r = block(ws, r, "Independent Route B found?", node["route_vs_result_analysis"]["was_an_independent_Route_B_found"], height=90)
r = block(ws, r, "Resolution", node["route_vs_result_analysis"]["resolution"], height=140)

# 250 Project State Graph
ws = new_sheet("250 REBUILD-003 State Graph")
write_title(ws, "PROJECT STATE GRAPH (16-run inventory)", "")
rows = [(r["id"], r["object"], r["status"]) for r in PSG["runs"]]
write_table(ws, 5, ["Run", "Object", "Status"], rows, widths=[45, 65, 55])

# 251 Universal Seed Registry
ws = new_sheet("251 REBUILD-004 Seed Registry")
write_title(ws, "UNIVERSAL SEED REGISTRY (updated with N=5)", "")
rows = [(c["name"], str(c["encoding_complexity_bits_N4"]), c["fixed_point_status"], c["canonicality"]) for c in SEEDREG["candidates"]]
write_table(ws, 5, ["Candidate", "K bits (N=4)", "Fixed-point status", "Canonicality"], rows, widths=[35, 15, 60, 40])

# 252 Compiler Implementation
ws = new_sheet("252 REBUILD-005 Compiler Impl")
write_title(ws, "EXECUTABLE COMPILER IMPLEMENTATION", "")
rows = [(m["file"], m["implements"], m["real_or_stub"]) for m in COMPREG["modules"]]
write_table(ws, 5, ["File", "Implements", "Real or stub?"], rows, widths=[30, 100, 40])
r = 10
r = block(ws, r, "Test results", "5/5 tests passing (~18s runtime): seed counts match prior runs, all survivors rigid, canonical-orientation-nilpotency theorem holds with zero exceptions, Gamma_1 non-idempotence counterexample reproduces, zero symmetric survivors through N=4 confirmed.", height=80)
r = block(ws, r, "Reproduce", COMPREG["reproducibility"], height=30)
r = block(ws, r, "Deliberately not built", COMPREG["what_was_deliberately_not_built"]["PASS-10_through_PASS-18_implementations"], height=140)

# 253 Final Closure
ws = new_sheet("253 REBUILD-006 Final Closure")
write_title(ws, "FINAL CLOSURE STATUS", "")
ws["A5"] = ("Kernel: complete and correctly reconfigured, now independently re-verified in executable "
            "code (5/5 tests passing).\n\n"
            "Seed family: existence proven through N=5 (1, 4, 23, 231); uniqueness disproven; growth "
            "accelerating (ratios 4.0, 5.75, 10.04), arguing against eventual convergence.\n\n"
            "Gauge group: reclassified via formal reconciliation (RECON-001) as EMPIRICAL INPUT for "
            "the tested route, OPEN for any untested route -- not falsified, not 'derived pending "
            "reconstruction.' Four independent source documents, not one, concur no derivation exists.\n\n"
            "Full compiler: OPEN, blocked on three explicitly identified, unresolved dependencies "
            "(lambda_gap, N-scaling/continuum behavior, gauge derivation route) -- none invented, "
            "none forced closed, none fabricated.\n\n"
            "Final Theorem: OPEN / PROVEN NON-IDENTIFIABLE (unchanged verdict, strengthened by the "
            "N=5 evidence against convergence).")
ws["A5"].font = BODY_FONT
ws["A5"].alignment = WRAP
ws.merge_cells("A5:F5")
ws.row_dimensions[5].height = 260

wb.save(OUT)
print("saved", OUT)
print("total sheets:", len(wb.sheetnames))
