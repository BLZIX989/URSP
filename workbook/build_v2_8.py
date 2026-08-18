"""
Build v2.8 from v2.7: preserves all 259 existing sheets, appends the
OPEN-024B (asymmetric NCG bridge) sheets.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import json

SRC = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v2.7.xlsx"
OUT = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v2.8.xlsx"
RECON = "/home/user/URSP/reconstruction"

def load(name):
    with open(f"{RECON}/{name}") as f:
        return json.load(f)

AXIOMS = load("ncg_asymmetric_axioms.json")
REPS = load("ncg_asymmetric_representations.json")
PROOFS = load("ncg_first_order_proofs.json")
EQUIV = load("ncg_representation_equivalence.json")
DFL = load("ncg_df_laplacian_relation.json")
KO = load("ncg_ko_dimension.json")
ALG = load("ncg_algebra_enumeration.json")
PD = load("ncg_poincare_duality.json")
ORI = load("ncg_orientability.json")
TOP = load("ncg_asymmetric_bridge.json")
with open(f"{RECON}/uoc_asymmetric_bridge_results.json") as f:
    SWEEP = json.load(f)

wb = openpyxl.load_workbook(SRC, data_only=False)

TITLE_FONT = Font(name="Arial", size=12, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")

def new_sheet(name):
    return wb.create_sheet(name)

def write_title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = WRAP

def write_table(ws, start_row, headers, rows, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + 1 + len(rows)

def block(ws, r, title, text, height=60):
    ws.cell(row=r, column=1, value=title).font = Font(bold=True)
    r += 1
    ws.cell(row=r, column=1, value=text)
    ws.cell(row=r, column=1).font = BODY_FONT
    ws.cell(row=r, column=1).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = height
    return r + 2


# 260 Overview
ws = new_sheet("260 ASYM-001 Overview")
write_title(ws, "OPEN-024B -- ASYMMETRIC NCG FUNCTOR BRIDGE", "")
r = block(ws, 5, "Result", "The obstruction is deeper than symmetric-vs-asymmetric: it is the ABELIAN algebra itself. All 4 predefined representation candidates (identity, conjugate, orientation_reversed, random_diagonal_perm), across all 259 seeds, both KO sign conventions: order-zero 100% pass, first-order 0% pass. PROVEN general theorem, not observed.", height=110)

# 261 Source Axioms
ws = new_sheet("261 ASYM-002 Source Axioms")
write_title(ws, "RECOVERED PROJECT NCG AXIOMS (verbatim from source sheets 152-157)", "")
rows = []
for k, v in AXIOMS["axioms"].items():
    rows.append((k, str(v)))
write_table(ws, 5, ["Axiom", "Statement/status"], rows, widths=[28, 130])

# 262 Representation Search Space
ws = new_sheet("262 ASYM-003 Search Space")
write_title(ws, "PREDEFINED REPRESENTATION SEARCH SPACE", "declared before evaluation")
rows = [(c["name"], c["formula"], c["classification"]) for c in REPS["search_space_declared_in_advance"]]
write_table(ws, 5, ["Candidate", "Formula", "Classification"], rows, widths=[24, 60, 45])
r = 11
r = block(ws, r, "Post-hoc validity audit", REPS["post_hoc_validity_audit"]["finding"], height=110)

# 263 Seed-Induced Representations
ws = new_sheet("263 ASYM-004 Seed-Induced")
write_title(ws, "SEED-INDUCED REPRESENTATION (orientation_reversed)", "")
r = block(ws, 5, "Definedness", REPS["orientation_reversed_undefined_cases"], height=140)

# 264 J Structure / KO
ws = new_sheet("264 ASYM-005 J and KO")
write_title(ws, "J STRUCTURE AND KO-DIMENSION", "")
r = block(ws, 5, "Result", KO["result"], height=80)
r = block(ws, r, "Uniquely selected by seed?", KO["is_it_uniquely_selected"], height=90)

# 265 Order Zero
ws = new_sheet("265 ASYM-006 Order Zero")
write_title(ws, "ORDER-ZERO -- 100% PASS", "")
write_table(ws, 5, ["N", "n_candidates"], [(n, SWEEP[n]["n_candidates"]) for n in SWEEP], widths=[10,15])

# 266 First Order
ws = new_sheet("266 ASYM-007 First Order")
write_title(ws, "FIRST-ORDER -- 0% PASS (primary result)", "")
rows = []
for N, data in SWEEP.items():
    for cand, agg in data["by_representation"].items():
        rows.append((N, cand, agg["tested"], agg["undefined"], agg["order_zero_pass"], agg["first_order_pass"]))
write_table(ws, 5, ["N", "Candidate", "Tested", "Undefined", "Order-zero pass", "First-order pass"], rows, widths=[6, 22, 10, 12, 16, 16])

# 267 D_F Construction
ws = new_sheet("267 ASYM-008 DF Construction")
write_title(ws, "D_F CONSTRUCTION", "unchanged from run 17 -- N_orient(R)+N_orient(R)^T, Hermitian, off-diagonal in grading basis")
r = block(ws, 5, "Note", "D_F is derived entirely from the seed's own canonical orientation, no free parameter, no fitting.", height=50)

# 268 D_F Spectrum (placeholder summary)
ws = new_sheet("268 ASYM-009 DF Spectrum")
write_title(ws, "D_F SPECTRUM", "")
r = block(ws, 5, "Note", "D_F's spectrum varies by seed (different edge patterns give different eigenvalues) -- this variation does not affect the first-order outcome (uniformly fails regardless), consistent with THM-ASYM-BRIDGE-OBSTRUCTION-001's proof that the obstruction is representation-theoretic (abelian algebra), not spectral.", height=90)

# 269 D_F^2 - L Comparison
ws = new_sheet("269 ASYM-010 DF2 vs L")
write_title(ws, "D_F^2 vs LAPLACIAN L -- TESTED, NEGATIVE", "")
r = block(ws, 5, "Result", DFL["result"], height=90)
r = block(ws, r, "Shifted relation (D_F^2=L+cI)?", DFL["no_shifted_relation_found"], height=60)
r = block(ws, r, "Status", DFL["status"], height=70)

# 270 KO Dimension (detail)
ws = new_sheet("270 ASYM-011 KO Dimension")
write_title(ws, "KO-DIMENSION DETAIL", "")
r = block(ws, 5, "Status", KO["status"], height=70)

# 271 Representation Equivalence
ws = new_sheet("271 ASYM-012 Rep Equivalence")
write_title(ws, "REPRESENTATION EQUIVALENCE / CANONICALITY", "")
r = block(ws, 5, "Result", EQUIV["result"], height=70)
r = block(ws, r, "Status", EQUIV["status"], height=90)

# 272 Seed Dependence
ws = new_sheet("272 ASYM-013 Seed Dependence")
write_title(ws, "SEED DEPENDENCE", "")
sd = TOP["seed_dependence_Section_15"]
r = block(ws, 5, "Question", sd["question"], height=40)
r = block(ws, r, "Finding", sd["finding"], height=140)

# 273 Poincare Duality
ws = new_sheet("273 ASYM-014 Poincare Duality")
write_title(ws, "POINCARE DUALITY -- DEFERRED", "")
r = block(ws, 5, "Status", PD["status"], height=90)
r = block(ws, r, "Context", PD["context"], height=70)

# 274 Orientability
ws = new_sheet("274 ASYM-015 Orientability")
write_title(ws, "ORIENTABILITY -- DEFERRED", "")
r = block(ws, 5, "Status", ORI["status"], height=70)
r = block(ws, r, "What IS established", ORI["what_IS_established"], height=90)
r = block(ws, r, "Context", ORI["context"], height=70)

# 275 Algebra Enumeration
ws = new_sheet("275 ASYM-016 Algebra Enum")
write_title(ws, "FINITE ALGEBRA SELECTION", "")
r = block(ws, 5, "What was tested", ALG["what_was_tested"], height=70)
r = block(ws, r, "Result", ALG["result"], height=90)
r = block(ws, r, "Does C(+)H(+)M3(C) emerge?", ALG["does_C_H_M3C_emerge"], height=90)
r = block(ws, r, "Status", ALG["status"], height=90)

# 276 Closure Matrix (status registry)
ws = new_sheet("276 ASYM-017 Closure Matrix")
write_title(ws, "OPEN-024B STATUS REGISTRY", "")
rows = [(k, v) for k, v in TOP["status_registry_Section_27"].items()]
write_table(ws, 5, ["Item", "Status"], rows, widths=[45, 110])

# 277 Proof Registry
ws = new_sheet("277 ASYM-018 Proof Registry")
write_title(ws, "THM-ASYM-BRIDGE-OBSTRUCTION-001 -- FULL PROOF", "")
r = block(ws, 5, "Statement", PROOFS["THM_ASYM_BRIDGE_OBSTRUCTION_001"]["statement"], height=140)
write_table(ws, r, ["Proof step"], [[s] for s in PROOFS["THM_ASYM_BRIDGE_OBSTRUCTION_001"]["proof"]], widths=[150])

# 278 Falsification Ledger
ws = new_sheet("278 ASYM-019 Falsification")
write_title(ws, "FALSIFICATION LEDGER", "")
write_table(ws, 5, ["Claim tested", "Result"], [
    ("Genuine representation asymmetry resolves the first-order obstruction", "FALSIFIED -- 0/2072 test instances pass, all 4 candidates, all 259 seeds"),
    ("The obstruction was specifically about pi_L=pi_R (not about the algebra being abelian)", "FALSIFIED -- disproven by the more general THM-ASYM-BRIDGE-OBSTRUCTION-001"),
    ("D_F^2 relates to the seed Laplacian L", "FALSIFIED -- exact non-equality on all 28 tested seeds, no shifted relation found either"),
    ("KO-dimension is uniquely selected by the seed", "FALSIFIED -- selected by an external doubling convention, not by R"),
], widths=[70, 90])

# 279 MDCL
ws = new_sheet("279 ASYM-020 MDCL")
write_title(ws, "CORRECTED MDCL FRAGMENT", "")
ws["A5"] = ("R (seed) -> bipartition, N_orient -> D_F, gamma_F [PROVEN general]\n"
            "                                  -> J_F (swap-conjugate) -> KO in {0,6} [PROVEN, seed-independent]\n"
            "                                  -> pi(A_F=C^N), any faithful representation\n"
            "                                  -> order-zero [PROVEN, always holds]\n"
            "                                  -> first-order [PROVEN, NEVER holds for abelian A_F]\n"
            "                                  -> STOP (THM-ASYM-BRIDGE-OBSTRUCTION-001)\n\n"
            "Poincare duality, orientability, finite algebra selection: DEFERRED, contingent on a "
            "non-abelian algebra bridge not yet built.")
ws["A5"].font = BODY_FONT
ws["A5"].alignment = WRAP
ws.merge_cells("A5:F5")
ws.row_dimensions[5].height = 180

# 280 Next Dependency
ws = new_sheet("280 ASYM-021 Next Dep")
write_title(ws, "EXACTLY ONE NEXT UNRESOLVED UPSTREAM DEPENDENCY", "")
ws["A4"] = ("Build and test a bridge variant using a genuinely NON-ABELIAN candidate algebra -- one "
            "containing at least one multi-dimensional irreducible representation, derived from the "
            "seed's own structure (e.g. a coarser block structure over the bipartition rather than a "
            "per-vertex diagonal) rather than an externally-motivated choice like H or M3(C). This is "
            "the precise, minimal condition THM-ASYM-BRIDGE-OBSTRUCTION-001's proof identifies as "
            "necessary for first-order to have any chance of closing with D_F != 0 -- not a guess, a "
            "direct consequence of the theorem. Not attempted this run, per the stop condition.")
ws["A4"].font = Font(name="Arial", size=12, bold=True)
ws["A4"].alignment = WRAP
ws.merge_cells("A4:F4")
ws.row_dimensions[4].height = 200

wb.save(OUT)
print("saved", OUT)
print("total sheets:", len(wb.sheetnames))
