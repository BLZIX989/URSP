"""
Build v1.5 from v1.4: preserves all 82 existing sheets, appends the C-004E
discrete-to-continuum correspondence closure sheets.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import json

SRC = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v1.4.xlsx"
OUT = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v1.5.xlsx"
RECON = "/home/user/URSP/reconstruction"

with open(f"{RECON}/c004e_registry.json") as f:
    R = json.load(f)
with open(f"{RECON}/c004e_kernel_check.json") as f:
    K = json.load(f)

wb = openpyxl.load_workbook(SRC, data_only=False)

TITLE_FONT = Font(name="Arial", size=12, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def new_sheet(name):
    return wb.create_sheet(name)


def write_title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = WRAP


def write_table(ws, start_row, headers, rows, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + 1 + len(rows)


# ---------------------------------------------------------------------
# 83 C-004E Proof Obligation
# ---------------------------------------------------------------------
ws = new_sheet("83 C-004E Proof Obligation")
write_title(ws, "C-004E -- EXACT PROOF OBLIGATION",
            "L*_discrete = -LD^{-1}  -->?  L*_continuum = -div(P grad log P_ss), under the project's own "
            "continuum-limit machinery (a sequence G_n -> M with scaling operators R_n). Frozen: everything "
            "at or above C-004D is unchanged by this run.")
ws["A5"] = R["exact_proof_obligation"]
ws["A5"].font = BODY_FONT
ws["A5"].alignment = WRAP
ws.merge_cells("A5:E5")
ws.row_dimensions[5].height = 40

# ---------------------------------------------------------------------
# 84 C-004E Convergence Machinery Search
# ---------------------------------------------------------------------
ws = new_sheet("84 C-004E Machinery Search")
write_title(ws, "SEARCH: DOES A G_n -> M CONVERGENCE MAP EXIST FOR THE THERMODYNAMIC BRANCH?", "")
write_table(ws, 4,
    ["Finding", "Detail"],
    [
        ("F1: Existing convergence machinery", R["findings"]["F1_existing_convergence_machinery"]["what_exists"]),
        ("F1: What it targets", R["findings"]["F1_existing_convergence_machinery"]["what_it_targets"]),
        ("F1: Operator it is built on", R["findings"]["F1_existing_convergence_machinery"]["operator_it_is_built_on"]),
        ("F1: Certification status", R["findings"]["F1_existing_convergence_machinery"]["certification_status"]),
        ("F2: Reference-measure ambiguity", R["findings"]["F2_reference_measure_ambiguity"]["finding"]),
        ("F2: Resolution", R["findings"]["F2_reference_measure_ambiguity"]["resolution"]),
        ("F2: Consequence", R["findings"]["F2_reference_measure_ambiguity"]["consequence"]),
        ("F3: No connection to the FP equation", R["findings"]["F3_no_connection_to_the_FP_equation"]["finding"]),
        ("F3: Consequence", R["findings"]["F3_no_connection_to_the_FP_equation"]["consequence"]),
    ], widths=[36, 110])

# ---------------------------------------------------------------------
# 85 C-004E Answer
# ---------------------------------------------------------------------
ws = new_sheet("85 C-004E Answer")
write_title(ws, "ANSWER TO THE EXACT PROOF OBLIGATION", "")
ws["A5"] = "Does the framework provide a sequence G_n -> M and scaling operators such that lim R_n(-L_n D_n^{-1}) = the continuum FP operator?"
ws["A5"].font = Font(bold=True)
ws["A5"].alignment = WRAP
ws.merge_cells("A5:E5")
ws.row_dimensions[5].height = 30
ws["A7"] = "NO."
ws["A7"].font = Font(bold=True, size=14, color="C00000")
ws["A9"] = R["answer_to_the_exact_proof_obligation"]["reasoning"]
ws["A9"].font = BODY_FONT
ws["A9"].alignment = WRAP
ws.merge_cells("A9:E9")
ws.row_dimensions[9].height = 110
ws["A11"] = R["answer_to_the_exact_proof_obligation"]["classification"]
ws["A11"].font = Font(bold=True)
ws["A11"].alignment = WRAP
ws.merge_cells("A11:E11")
ws.row_dimensions[11].height = 55

# ---------------------------------------------------------------------
# 86 THM-GEN-DISTINCT-001
# ---------------------------------------------------------------------
ws = new_sheet("86 THM-GEN-DISTINCT-001")
write_title(ws, "CANONICAL THEOREM: SPECTRAL/THERMODYNAMIC GENERATOR DISTINCTION",
            "Promoted from the C-004D finding to a general, ARBS-independent structural theorem, per this run's directive.")
ws["A5"] = R["general_theorem_promoted_to_canonical"]["statement"]
ws["A5"].font = BODY_FONT
ws["A5"].alignment = WRAP
ws.merge_cells("A5:E5")
ws.row_dimensions[5].height = 150
r = 12
ws.cell(row=r, column=1, value="Status").font = Font(bold=True)
ws.cell(row=r, column=2, value=R["general_theorem_promoted_to_canonical"]["status"])
r += 1
ws.cell(row=r, column=1, value="Scope").font = Font(bold=True)
ws.cell(row=r, column=2, value=R["general_theorem_promoted_to_canonical"]["scope"])
ws.cell(row=r, column=2).alignment = WRAP
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
ws.row_dimensions[r].height = 40
r += 2
ws.cell(row=r, column=1, value="Numerical confirmation: ||L*d|| (should be 0 iff d in ker(L), i.e. iff G is regular)").font = Font(bold=True)
r += 1
rows = [(k, v["regular"], v["||L*d||"], v["d_in_ker(L)"]) for k, v in K.items()]
write_table(ws, r, ["Shell", "Regular?", "||L*d||", "d in ker(L)?"], rows, widths=[8, 10, 14, 14])

# ---------------------------------------------------------------------
# 87 C-004E Closure
# ---------------------------------------------------------------------
ws = new_sheet("87 C-004E Closure")
write_title(ws, "C-004E FINAL CLOSURE", "")
write_table(ws, 4,
    ["Object", "Status"],
    [
        ("C-004E (discrete-to-continuum FP correspondence)", "OPEN -- precise theorem-level obstruction identified: no convergence machinery connecting any thermodynamic-branch discretization to the continuum FP equation exists anywhere in the corpus"),
        ("THM-GEN-DISTINCT-001", "PROVEN; recommended for canonical promotion (general, ARBS-independent result)"),
        ("C-004D (generator reconciliation, discrete level)", "Unaffected, still PARTIALLY CLOSED"),
        ("C-004C (sigma_k identifiability)", "Not reopened; sigma_k remains PROVEN NON-IDENTIFIABLE"),
        ("G*", "UNRESOLVED (unchanged)"),
    ], widths=[46, 100])

# ---------------------------------------------------------------------
# 88 C-004E Next Dependency
# ---------------------------------------------------------------------
ws = new_sheet("88 C-004E Next Dependency")
write_title(ws, "EXACTLY ONE NEXT UNRESOLVED UPSTREAM DEPENDENCY", "")
ws["A4"] = ("Construct (or formally register as a new, explicitly named open bridge in the source registry) a "
            "discrete-to-continuum convergence theorem for the THERMODYNAMIC branch specifically: a sequence "
            "G_n -> M, scaling operators R_n, and a proof (Mosco convergence, Gromov-Hausdorff, or an equivalent "
            "route already admissible in this project) showing lim R_n(-L_n D_n^{-1}) converges to some explicit "
            "continuum operator -- and only then determine whether that limit is the specific PDE "
            "L*P=-div(P grad log P_ss) already asserted in source, or a different continuum object. "
            "This must be built in the degree-weighted L^2(V,d) Hilbert space (the space in which -D^{-1}L is "
            "self-adjoint, proven in C-004D), NOT simply re-used from the existing counting-measure RF-001..005 "
            "apparatus, which was verified this run to be built on a different reference measure entirely. "
            "Until this exists, sigma_k's generator has a derived discrete form (C-004D) but no verified continuum "
            "meaning, and no further attempt to pin down P(0)/evaluation time for C-004C should proceed.")
ws["A4"].font = Font(name="Arial", size=12, bold=True)
ws["A4"].alignment = WRAP
ws.merge_cells("A4:E4")
ws.row_dimensions[4].height = 160

wb.save(OUT)
print("saved", OUT)
print("total sheets:", len(wb.sheetnames))
