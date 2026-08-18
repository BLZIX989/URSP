"""
Build v1.3 from v1.2: preserves all 62 existing sheets, appends the C-004C
entropy-production recovery closure sheets.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import json

SRC = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v1.2.xlsx"
OUT = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v1.3.xlsx"
RECON = "/home/user/URSP/reconstruction"

with open(f"{RECON}/c004c_sigma_registry.json") as f:
    REG = json.load(f)
with open(f"{RECON}/sigma_nonidentifiability_demo.json") as f:
    DEMO = json.load(f)

wb = openpyxl.load_workbook(SRC, data_only=False)

TITLE_FONT = Font(name="Arial", size=12, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def new_sheet(name):
    return wb.create_sheet(name)


def write_title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = WRAP


def write_table(ws, start_row, headers, rows, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + 1 + len(rows)


# ---------------------------------------------------------------------
# 63 C-004C sigma Search
# ---------------------------------------------------------------------
ws = new_sheet("63 C-004C sigma Search")
write_title(ws, "C-004C -- CORPUS SEARCH FOR ENTROPY-PRODUCTION sigma",
            "Search executed across all 9 source workbooks + whitepaper for sigma / entropy "
            "production / H-theorem / dS/dt / Clausius-Duhem / Onsager / dissipation and "
            "equivalent unnamed definitions (Section 6).")
write_table(ws, 5,
    ["Found", "Location", "Content"],
    [
        ("YES", "BRIDGE B-006 (CERTIFIED)", "sigma = -dF/dt >= 0, F[P]=int P log(P/P_ss) dmu, H-theorem for Fokker-Planck"),
        ("YES", "DER-TD-004 / TD-004 (CERTIFIED)", "sigma = sum_i J_i X_i >= 0, Onsager flux-force local entropy production"),
        ("YES", "DER-TD-005 / DER-ORG-007 / DER-OPEN-003 (all OPEN)", "Proposed alternate coupling Pi_O=1-exp(-sigma*delta_spec/lambda_c); Pi_O=f(sigma,C,Spec(L)) well-definedness itself registered OPEN"),
        ("YES (negative finding)", "H_FP hyperedge registry", "'Stationary measure requires generator, initial distribution, and domain' -- initial distribution flagged as a required, unsupplied input"),
        ("YES (only precedent)", "BRIDGE B-007 / C-008", "a(0)=e_1 used for K3,3 (a DIFFERENT test graph, not any ARBS shell) as an ad hoc illustrative choice, not a general rule"),
        ("NO", "Any registry entry for a per-shell numeric sigma_k, k=0..4", "Not found anywhere in the 9-workbook + whitepaper corpus"),
    ], widths=[22, 40, 80])

# ---------------------------------------------------------------------
# 64 C-004C sigma Registry
# ---------------------------------------------------------------------
ws = new_sheet("64 C-004C sigma Registry")
write_title(ws, "sigma CANDIDATE PROVENANCE TABLE", "")
headers = ["Candidate ID", "Name", "Exact equation", "IC dependence", "Time dependence",
           "Target dependence", "Status"]
rows = [(c["candidate_id"], c["name"], c["exact_equation"],
         c.get("initial_condition_dependence", "n/a"), c.get("time_dependence", "n/a"),
         c.get("target_dependence", "n/a"), c["status"])
        for c in REG["candidates"]]
write_table(ws, 4, headers, rows, widths=[16, 34, 50, 34, 34, 20, 46])

# ---------------------------------------------------------------------
# 65 C-004C H-Theorem Audit
# ---------------------------------------------------------------------
ws = new_sheet("65 C-004C H-Theorem Audit")
write_title(ws, "H-THEOREM ANALYSIS (Section 9)", "")
write_table(ws, 4,
    ["Question", "Answer"],
    [
        ("What is S_k / F[P]?", "F[P] = integral P log(P/P_ss) dmu -- KL divergence of P from the stationary measure P_ss (CERTIFIED, B-006)"),
        ("What is the state variable?", "A time-dependent probability distribution P(t) over the graph's vertex set"),
        ("What evolution equation determines the state?", "d_t P = L* P, a Fokker-Planck-type generator; NOT identified with certainty against the certified combinatorial Laplacian L=D-A (see GAP-GENERATOR-001, sheet 66)"),
        ("What boundary conditions exist?", "None specified for finite ARBS shells"),
        ("What initial conditions exist?", "None fixed canonically; only precedent is an arbitrary a(0)=e_1 on an unrelated test graph (K3,3), sheet 63"),
        ("Is sigma_k uniquely determined by G_k?", "NO -- depends on P(0) and t (and on which generator), none of which G_k alone fixes"),
        ("Is sigma_k constant?", "NO -- sigma(t) is a function of time along the trajectory, generically decreasing to 0 as P(t)->P_ss"),
        ("Is sigma_k time-dependent?", "YES"),
        ("Is there a canonical evaluation time?", "NOT FOUND in source"),
        ("Is there a canonical shell normalization?", "NOT FOUND in source"),
    ], widths=[46, 90])

# ---------------------------------------------------------------------
# 66 C-004C Initial Condition Audit
# ---------------------------------------------------------------------
ws = new_sheet("66 C-004C IC Audit")
write_title(ws, "INITIAL-CONDITION / GENERATOR AUDIT (Section 13)", "")
r = write_table(ws, 4,
    ["Required datum", "Canonically fixed by ARBS/UOC?", "Detail"],
    [
        ("P(0) / rho(0) / initial distribution", "NO",
         "H_FP hyperedge registry explicitly lists it as required-but-unsupplied; only precedent (a(0)=e_1) is example-specific to a different graph"),
        ("Evaluation time t (or reduction rule: asymptotic, integral, etc.)", "NO", "Not specified anywhere"),
        ("Generator L* (Fokker-Planck operator)", "AMBIGUOUS",
         "BRIDGE B-006 implies a generator whose stationary distribution is the degree distribution "
         "d_i/(2|E|); the certified ARBS Laplacian L=D-A (used throughout the rest of the project) "
         "has the UNIFORM distribution as its stationary state for a connected graph -- these differ "
         "for any non-regular graph, and G0-G4 are not regular. No registry entry reconciles the two."),
    ], widths=[46, 26, 90])
r += 1
ws.cell(row=r, column=1, value="OUTCOME: C -- entropy production is NOT uniquely determined by the current substrate.")
ws.cell(row=r, column=1).font = Font(bold=True, color="C00000")
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)

# ---------------------------------------------------------------------
# 67 C-004C Dimensional Audit
# ---------------------------------------------------------------------
ws = new_sheet("67 C-004C Dimensional Audit")
write_title(ws, "DIMENSIONAL CONSISTENCY AUDIT (Section 14)", "")
write_table(ws, 4,
    ["Check", "Result"],
    [
        ("Units of delta_spec, lambda_c", "Both are eigenvalues of the graph Laplacian L -- same units by construction, so (delta_spec-lambda_c) is well-formed"),
        ("Units of sigma = -dF/dt", "F[P] is a KL divergence (dimensionless, in nats). t is paired with L's eigenvalues via exp(-t*lambda) (t has units of 1/[lambda]). "
         "Therefore d(dimensionless)/d(1/[lambda]) has units of [lambda] -- matching delta_spec and lambda_c."),
        ("Verdict", "PASSES structurally -- sigma's units are consistent with delta_spec/lambda_c IF the same generator/time convention is used throughout. "
         "This says nothing about sigma's VALUE, which remains undetermined (sheet 66)."),
    ], widths=[30, 100])

# ---------------------------------------------------------------------
# 68 C-004C Target Independence
# ---------------------------------------------------------------------
ws = new_sheet("68 C-004C Target Independence")
write_title(ws, "TARGET-INDEPENDENCE AUDIT (Section 12)", "")
write_table(ws, 4,
    ["Candidate", "d(sigma)/d(N_H, H_F, D_F, J_F, gamma_F, Y_u, Y_d, Y_e, Y_nu, V_CKM, U_PMNS, Xi_obs)"],
    [
        ("SIGMA-CAND-001 (H-theorem)", "= 0 for every listed quantity -- the formula contains no downstream object, no fitted constant"),
        ("SIGMA-CAND-002 (Onsager)", "= 0 for every listed quantity by the same argument"),
        ("SIGMA-CAND-003 (proposed alt. coupling)", "= 0 structurally, but candidate itself is OPEN/PROPOSED in source, not admitted"),
    ], widths=[36, 100])
r = 9
ws.cell(row=r, column=1, value="All candidates PASS target-independence. This audit does not resolve non-identifiability; it only confirms the free data (P(0), t, generator) are not secretly encoding a downstream target.")
ws.cell(row=r, column=1).font = NOTE_FONT
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
ws.row_dimensions[r].height = 30

# ---------------------------------------------------------------------
# 69 C-004C Calculations
# ---------------------------------------------------------------------
ws = new_sheet("69 C-004C Calculations")
write_title(ws, "SIGN RESULT AND PARAMETER-DEPENDENCE DEMONSTRATION",
            "Section 16/17. Illustrative demonstration only -- no choice here is canonical.")
r = write_table(ws, 5,
    ["Result", "Value"],
    [
        ("sigma_k > 0 (generic case, any admissible generator/IC)", "PROVEN (H-theorem, BRIDGE B-006)"),
        ("delta_spec(G_k) - lambda_c,k, all k", "< 0 for every k=0..4 (established prior run)"),
        ("Pi_0(G_k), all k", "< 0 for every k=0..4 (sign only; magnitude undetermined)"),
    ], widths=[50, 50])
r += 2
ws.cell(row=r, column=1, value="Illustrative argmax_k Pi_0(G_k) under 6 sampled (P(0), t) combinations (certified heat semigroup exp(-tL), uniform P_ss):").font = Font(bold=True)
r += 1
combo_rows = [(combo, v["argmax_shell"], json.dumps(v["values"])) for combo, v in DEMO["combo_argmax"].items()]
r = write_table(ws, r, ["(P(0), t) combination", "argmax shell", "Pi0 values (all shells)"], combo_rows, widths=[30, 14, 90])
r += 1
ws.cell(row=r, column=1,
        value="Result: argmax takes 3 different values (G0, G1, G2) across 6 illustrative samples. This is a "
              "concrete existence proof of parameter-dependence -- confirms Section 17's 'Parameter-dependent "
              "ranking' outcome. It is not a claim about the true ranking under any canonical choice, because "
              "no such choice exists in the source material.")
ws.cell(row=r, column=1).font = NOTE_FONT
ws.cell(row=r, column=1).alignment = WRAP
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
ws.row_dimensions[r].height = 55

# ---------------------------------------------------------------------
# 70 C-004C Closure
# ---------------------------------------------------------------------
ws = new_sheet("70 C-004C Closure")
write_title(ws, "C-004C FINAL CLOSURE", "")
r = write_table(ws, 4,
    ["Object", "Status"],
    [
        ("sigma_k (specific per-shell number)", "PROVEN NON-IDENTIFIABLE"),
        ("sigma_k (sign)", "PROVEN SIGN-ONLY: sigma_k > 0 (generic case)"),
        ("Pi_0(G_k) magnitude/ranking", "OPEN"),
        ("Pi_0(G_k) sign", "PROVEN SIGN-ONLY: Pi_0(G_k) < 0 for all k=0..4"),
        ("OPEN-020E.2", "REMAINS OPEN -- argmax_k Pi_0(G_k) not evaluable; demonstrated parameter-dependent"),
        ("G*", "UNRESOLVED (unchanged)"),
        ("GAP-GENERATOR-001 (new)", "OPEN -- generator ambiguity between BRIDGE B-006's degree-distribution P_ss "
         "and the certified combinatorial Laplacian's uniform stationary state"),
    ], widths=[36, 90])

# ---------------------------------------------------------------------
# 71 OPEN-020E.2 Updated Audit
# ---------------------------------------------------------------------
ws = new_sheet("71 OPEN-020E.2 Updated")
write_title(ws, "OPEN-020E.2 UPDATED AUDIT (post C-004C)", "")
write_table(ws, 4,
    ["Component", "Status this run"],
    [
        ("delta_spec(G_k)", "CALCULATED (prior run: = lambda_1(L_k) for each shell)"),
        ("lambda_c,k", "CALCULATED (prior run: = 1/t_H,k)"),
        ("sigma_k", "PROVEN NON-IDENTIFIABLE (this run) -- sign only, sigma_k > 0"),
        ("Pi_0(G_k)", "Sign derived: < 0 for all k. Magnitude/ranking: OPEN"),
        ("argmax_k Pi_0(G_k) = 3 ?", "NOT TESTABLE -- outcome D persists: cannot be evaluated because sigma_k is undetermined"),
        ("G* = G3 promotion", "NOT PERFORMED"),
    ], widths=[30, 90])

wb.save(OUT)
print("saved", OUT)
print("total sheets:", len(wb.sheetnames))
