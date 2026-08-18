"""
Build v2.7 from v2.6: preserves all 253 existing sheets, appends the
SEMANTIC-FUNCTOR-BRIDGE-001 sheets.
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import json

SRC = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v2.6.xlsx"
OUT = "/home/user/URSP/workbook/Universal_Rosetta_Stone_MASTER_CALCULATION_WORKBOOK_v2.7.xlsx"
RECON = "/home/user/URSP/reconstruction"

def load(name):
    with open(f"{RECON}/{name}") as f:
        return json.load(f)

BRIDGE = load("uoc_semantic_functor_bridge.json")
RECON2 = load("reconciliation_RECON002_ncg_workbook.json")

wb = openpyxl.load_workbook(SRC, data_only=False)

TITLE_FONT = Font(name="Arial", size=12, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BODY_FONT = Font(name="Arial", size=10)
NOTE_FONT = Font(name="Arial", size=9, italic=True)
WRAP = Alignment(wrap_text=True, vertical="top")

def new_sheet(name):
    return wb.create_sheet(name)

def write_title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = WRAP

def write_table(ws, start_row, headers, rows, widths=None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BODY_FONT
            c.alignment = WRAP
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + 1 + len(rows)

def block(ws, r, title, text, height=60):
    ws.cell(row=r, column=1, value=title).font = Font(bold=True)
    r += 1
    ws.cell(row=r, column=1, value=text)
    ws.cell(row=r, column=1).font = BODY_FONT
    ws.cell(row=r, column=1).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = height
    return r + 2


# 254 NCG Workbook Discovery
ws = new_sheet("254 BRIDGE-001 NCG Discovery")
write_title(ws, "NCG WORKBOOK DISCOVERY (RECON-002)", "156-sheet dedicated workbook, never opened in runs 1-16")
r = block(ws, 5, "What it contains", RECON2["what_the_workbook_actually_contains"]["summary"], height=90)
r = block(ws, r, "V2-CALC-001 principal result", RECON2["what_the_workbook_actually_contains"]["V2_CALC_001_principal_result"], height=60)
write_table(ws, r, ["Sheet 09_NCG_SM_CLOSURE self-reported status"], [[x] for x in RECON2["what_the_workbook_actually_contains"]["explicit_self_reported_status_09_NCG_SM_CLOSURE_sheet"]], widths=[150])

# 255 ARBS Retirement Investigation
ws = new_sheet("255 BRIDGE-002 ARBS Retirement")
write_title(ws, "WHY ARBS IS CALLED 'RETIRED' -- INVESTIGATION RESULT", "")
r = block(ws, 5, "Conclusion", "AUD-010 is a scoping decision by a separate, non-engaging research thread, not a technical refutation. Zero mentions anywhere in the 156-sheet NCG workbook of tilde_G_k, Gamma_infty, R*, bipartite, shell, K_{2^k,2^k}, the SEF/SIT specification, or the DTC/Delta/Psi_t+1 grammar notation. The retirement line itself carries no calculation-ID citation, unlike every other status claim in the same workbook. Execution dates (2026-08-17, one day before this session's own runs) support a separate/parallel-thread explanation.", height=140)

# 256 Semantic Functor Bridge Construction
ws = new_sheet("256 BRIDGE-003 Construction")
write_title(ws, "SEMANTIC FUNCTOR BRIDGE -- CONSTRUCTION", "Track A (relational seed) -> Track B (finite NCG spectral triple)")
r = block(ws, 5, "Construction", BRIDGE["construction"], height=140)
r = block(ws, r, "Target independence", BRIDGE["target_independence"], height=50)

# 257 Bridge Results
ws = new_sheet("257 BRIDGE-004 Results")
write_title(ws, "BRIDGE TEST RESULTS -- exhaustive over all 259 candidates, N=2..5", "")
res = BRIDGE["results_exhaustive_over_all_259_candidates_N_2_through_5"]
write_table(ws, 5, ["Axiom", "Result"], [
    ("Phase 1: grading + Hermiticity + off-diagonality", res["phase1_structural_axioms"]),
    ("Phase 2: KO-dimension class", res["phase2_KO_dimension"]),
    ("Phase 3: order-zero (generic abelian algebra)", res["phase3_order_zero"]),
    ("Phase 3: first-order (generic abelian algebra)", res["phase3_first_order"]),
], widths=[45, 110])

# 258 Bridge Theorems
ws = new_sheet("258 BRIDGE-005 Theorems")
write_title(ws, "PROVEN GENERAL THEOREMS (not empirical coincidences)", "")
th = BRIDGE["these_results_are_PROVEN_general_theorems_not_empirical_coincidences"]
r = block(ws, 5, "THM-BRIDGE-ORDER-ZERO-001", th["THM-BRIDGE-ORDER-ZERO-001"], height=110)
r = block(ws, r, "THM-BRIDGE-FIRST-ORDER-OBSTRUCTION-001", th["THM-BRIDGE-FIRST-ORDER-OBSTRUCTION-001"], height=140)

# 259 Bridge Interpretation
ws = new_sheet("259 BRIDGE-006 Interpretation")
write_title(ws, "HONEST INTERPRETATION", "")
hi = BRIDGE["honest_interpretation"]
r = block(ws, 5, "What this establishes", hi["what_this_establishes"], height=80)
r = block(ws, r, "What this does NOT establish", hi["what_this_does_NOT_establish"], height=80)
r = block(ws, r, "Why the obstruction is structurally expected", hi["why_the_obstruction_is_structurally_expected"], height=110)
r = block(ws, r, "Next dependency", hi["next_dependency"], height=90)

wb.save(OUT)
print("saved", OUT)
print("total sheets:", len(wb.sheetnames))
